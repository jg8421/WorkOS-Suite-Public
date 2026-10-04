"""Registered-root file access, bounded previews and recoverable operations."""
from __future__ import annotations
from datetime import datetime
import csv
import hashlib
from html.parser import HTMLParser
import json
import io
import math
import mimetypes
import os
from pathlib import Path
import re
import shutil
import sqlite3
import stat
import threading
import time
import uuid
import xml.etree.ElementTree as ET
import zipfile

MAX_PREVIEW_BYTES=20_000_000
MAX_TREE_BYTES=512_000_000
MAX_TREE_ENTRIES=10_000
CLOUD_TAGS=frozenset(0x9000001A+(n<<12) for n in range(16))


class FileConflict(ValueError):
    status_code=409


def _linked(path):
    info=path.lstat()
    if stat.S_ISLNK(info.st_mode):return True
    if not getattr(info,'st_file_attributes',0)&0x400:return False
    tag=getattr(info,'st_reparse_tag',0)
    # Microsoft Cloud Files placeholder tags are not name-surrogate path links.
    return not (isinstance(tag,int) and not isinstance(tag,bool) and not tag&0x20000000 and tag in CLOUD_TAGS)


def _relative(value):
    if not isinstance(value,str) or len(value)>4096 or '\0' in value:raise ValueError('文件相对路径无效')
    value=value.replace('\\','/')
    if value.startswith('/') or ':' in value or any(p in ('.','..') for p in value.split('/')):raise ValueError('路径不能包含上级跳转或绝对地址')
    for part in value.split('/'):
        if re.search(r'[<>"|?*\x00-\x1f]',part) or part.endswith((' ','.')) or part.split('.',1)[0].upper() in {'CON','PRN','AUX','NUL',*(f'COM{i}' for i in range(1,10)),*(f'LPT{i}' for i in range(1,10))}:
            raise ValueError('文件名称包含不受支持的字符或系统保留名称')
    return '/'.join(p for p in value.split('/') if p)


def _plain(base,target,must_exist=True):
    base=Path(base);target=Path(target)
    for ancestor in reversed((base,*base.parents)):
        try:
            if _linked(ancestor):raise ValueError('登记目录路径含链接或不受支持的重解析点')
        except FileNotFoundError:pass
    base=base.resolve()
    try:target.relative_to(base)
    except ValueError:raise ValueError('文件超出已登记目录') from None
    current=base
    for part in target.relative_to(base).parts:
        current=current/part
        try:
            if _linked(current):raise ValueError('文件路径含链接或不受支持的重解析点')
        except FileNotFoundError:pass
    canonical=target.resolve(strict=False)
    try:canonical.relative_to(base)
    except ValueError:raise ValueError('文件超出已登记目录') from None
    if must_exist and not canonical.exists():raise FileNotFoundError('文件不存在')
    return canonical


def _tree(path):
    entries=[];total=0;pending=[path]
    while pending:
        item=pending.pop()
        if _linked(item):raise ValueError('文件夹包含链接，未执行操作')
        info=item.stat()
        if not stat.S_ISREG(info.st_mode) and not stat.S_ISDIR(info.st_mode):raise ValueError('不支持此文件类型')
        total+=info.st_size if item.is_file() else 0
        entries.append((str(item.relative_to(path)) if item!=path else '',item.is_dir(),info.st_size if item.is_file() else 0,info.st_mtime_ns if item.is_file() else 0))
        if len(entries)>MAX_TREE_ENTRIES or total>MAX_TREE_BYTES:raise ValueError('操作内容过大，请分批处理')
        if item.is_dir():
            for child in item.iterdir():
                pending.append(child)
                if len(pending)+len(entries)>MAX_TREE_ENTRIES:raise ValueError('操作内容过大，请分批处理')
    return sorted(entries)


def _fingerprint(path):
    records=_tree(path);digest=hashlib.sha256(json.dumps(records,sort_keys=True).encode())
    for relative,is_dir,*_ in records:
        if is_dir:continue
        item=path/relative if relative else path
        with item.open('rb') as stream:
            for chunk in iter(lambda:stream.read(1024*1024),b''):digest.update(chunk)
    return digest.hexdigest()


def _json(value):return json.dumps(value,ensure_ascii=False,sort_keys=True,allow_nan=False,separators=(',',':'))


class FileService:
    def __init__(self,data_dir,default_roots=None):
        self.data_dir=Path(data_dir).resolve();self.data_dir.mkdir(parents=True,exist_ok=True)
        self.recycle=self.data_dir/'recycle';self.recycle.mkdir(exist_ok=True)
        self._lock=threading.RLock()
        self._db=sqlite3.connect(self.data_dir/'files.sqlite3',check_same_thread=False)
        self._db.execute('PRAGMA journal_mode=WAL')
        self._db.execute('CREATE TABLE IF NOT EXISTS roots(id TEXT PRIMARY KEY,path TEXT UNIQUE,label TEXT)')
        self._db.execute('CREATE TABLE IF NOT EXISTS receipts(id TEXT PRIMARY KEY,request_id TEXT UNIQUE,payload TEXT NOT NULL,receipt TEXT NOT NULL)')
        self._db.commit()
        for root in default_roots or []:
            if Path(root).is_dir():self.add_root(root,Path(root).name)

    def close(self):
        with self._lock:self._db.close()

    def roots(self):
        with self._lock:return [dict(zip(('id','path','label'),row)) for row in self._db.execute('SELECT id,path,label FROM roots ORDER BY label,id')]

    def add_root(self,path,label=''):
        if not isinstance(path,(str,Path)) or not Path(path).is_absolute():raise ValueError('请登记一个完整的文件夹路径')
        original=Path(path)
        if '..' in original.parts:raise ValueError('目录地址不能包含上级跳转')
        canonical=original.resolve(strict=True)
        if not canonical.is_dir():raise ValueError('只能登记文件夹')
        # An explicit root registration authorizes this canonical directory, not its outside descendants.
        key='root_'+hashlib.sha256(os.path.normcase(str(canonical)).encode()).hexdigest()[:20]
        if not isinstance(label,str) or len(label)>100:raise ValueError('目录名称无效')
        with self._lock:
            self._db.execute('INSERT INTO roots VALUES(?,?,?) ON CONFLICT(id) DO UPDATE SET label=excluded.label',(key,str(canonical),label.strip() or canonical.name or str(canonical)))
            self._db.commit()
        return {'id':key,'path':str(canonical),'label':label.strip() or canonical.name or str(canonical)}

    def _root(self,root_id):
        if not isinstance(root_id,str) or len(root_id)>100:raise ValueError('目录标识无效')
        with self._lock:row=self._db.execute('SELECT path FROM roots WHERE id=?',(root_id,)).fetchone()
        if not row:raise ValueError('目录尚未登记')
        return Path(row[0])

    def resolve(self,root_id,path='',must_exist=True):
        relative=_relative(path);root=self._root(root_id)
        return _plain(root,root/relative,must_exist)

    def list(self,root_id,path='',q=''):
        relative=_relative(path);folder=self.resolve(root_id,relative)
        if not folder.is_dir():raise ValueError('请选择文件夹')
        if not isinstance(q,str) or len(q)>200:raise ValueError('搜索文字无效')
        entries=[]
        for item in folder.iterdir():
            if q.casefold() not in item.name.casefold():continue
            try:
                _plain(self._root(root_id),item);info=item.stat()
                entries.append({'name':item.name,'path':f'{relative}/{item.name}'.lstrip('/'),'is_dir':item.is_dir(),
                    'size':info.st_size if item.is_file() else 0,'modified':info.st_mtime,'mime':mimetypes.guess_type(item.name)[0] or 'application/octet-stream'})
                if len(entries)>=5000:break
            except (ValueError,OSError):continue
        return {'root_id':root_id,'path':relative,'parent':relative.rsplit('/',1)[0] if '/' in relative else '',
                'entries':sorted(entries,key=lambda e:(not e['is_dir'],e['name'].casefold()))}

    @staticmethod
    def _zip(path):
        archive=zipfile.ZipFile(path);infos=archive.infolist()
        if len(infos)>5000 or len({i.filename for i in infos})!=len(infos):archive.close();raise ValueError('压缩包条目过多或重复')
        total=0
        for info in infos:
            total+=info.file_size
            if info.file_size>MAX_PREVIEW_BYTES or total>50_000_000 or info.file_size>max(100_000,info.compress_size*100):
                archive.close();raise ValueError('压缩内容过大，未展开预览')
        return archive

    @staticmethod
    def _xml(archive,name):
        info=archive.getinfo(name)
        if info.file_size>2_000_000:raise ValueError('文档结构过大，未展开预览')
        raw=archive.read(name)
        if b'<!DOCTYPE' in raw.upper() or b'<!ENTITY' in raw.upper():raise ValueError('文档包含不受支持的 XML 定义')
        return ET.fromstring(raw)

    def preview(self,root_id,path):
        target=self.resolve(root_id,path)
        if not target.is_file():raise ValueError('请选择文件进行预览')
        if target.stat().st_size>MAX_PREVIEW_BYTES:raise ValueError('文件超过 20 MB，请下载原文件查看')
        suffix=target.suffix.casefold();base={'content':'','rows':[],'columns':[]}
        try:
            if suffix in ('.docx','.pptx','.xlsx','.zip'):
                with self._zip(target) as archive:
                    if suffix=='.zip':
                        return {**base,'type':'archive','rows':[[i.filename,i.file_size] for i in archive.infolist()[:200]],'columns':['文件','字节']}
                    if suffix=='.xlsx':
                        import openpyxl
                        book=openpyxl.load_workbook(target,read_only=True,data_only=True,keep_links=False)
                        try:
                            sheet=book.worksheets[0];rows=[]
                            for row in sheet.iter_rows(max_row=min(sheet.max_row or 200,200),max_col=min(sheet.max_column or 30,30),values_only=True):
                                rows.append([self._cell(value) for value in row])
                            return {**base,'type':'table','content':f'工作表：{sheet.title}；最多 200 行、30 列，仅预览缓存值，不重新计算公式。',
                                    'rows':rows,'columns':[openpyxl.utils.get_column_letter(i+1) for i in range(len(rows[0]) if rows else 0)]}
                        finally:book.close()
                    names=['word/document.xml'] if suffix=='.docx' else sorted((n for n in archive.namelist() if re.fullmatch(r'ppt/slides/slide\d+\.xml',n)),key=lambda n:int(re.search(r'(\d+)\.xml$',n).group(1)))[:100]
                    chunks=[]
                    for name in names:
                        tree=self._xml(archive,name)
                        chunks.extend(node.text or '' for node in tree.iter() if node.tag.rsplit('}',1)[-1]=='t')
                        if sum(map(len,chunks))>100_000:break
                    return {**base,'type':'text','content':'\n'.join(chunks)[:100_000]}
            if suffix=='.pdf':
                from pypdf import PdfReader
                reader=PdfReader(str(target))
                if reader.is_encrypted:raise ValueError('加密 PDF 需在原应用中查看')
                chunks=[]
                for index,page in enumerate(reader.pages[:30]):
                    chunks.append(f'第 {index+1} 页\n'+(page.extract_text() or '')[:100_000])
                    if sum(map(len,chunks))>=100_000:break
                return {**base,'type':'text','content':'\n\n'.join(chunks)[:100_000]}
            if suffix in ('.txt','.md','.csv','.tsv','.json','.log','.xml','.html','.htm','.py','.js','.css'):
                with target.open('rb') as stream:raw=stream.read(400_000)
                try:text=raw.decode('utf-8-sig')
                except UnicodeDecodeError:
                    try:text=raw.decode('gb18030')
                    except UnicodeDecodeError:text=raw.decode('utf-8',errors='replace')
                if suffix in ('.html','.htm'):
                    parser=_HTMLText();parser.feed(text);text='\n'.join(parser.parts)
                if suffix in ('.csv','.tsv'):
                    reader=csv.reader(io.StringIO(text),delimiter='\t' if suffix=='.tsv' else ',')
                    rows=[]
                    for index,row in enumerate(reader):
                        if index>=200:break
                        rows.append([value[:1000] for value in row[:30]])
                    width=max(map(len,rows),default=0)
                    return {**base,'type':'table','rows':rows,'columns':[str(i+1) for i in range(width)],'truncated':len(raw)>=400_000 or len(rows)>=200}
                return {**base,'type':'text','content':text[:100_000]}
            return {**base,'type':'binary','content':'此格式请下载原文件查看。'}
        except ImportError:raise ValueError('预览依赖尚未安装，请使用完整套件') from None
        except (zipfile.BadZipFile,ET.ParseError,KeyError):raise ValueError('文档结构无效，未展开预览') from None
        except ValueError:raise
        except Exception:raise ValueError('此文件暂无法预览，请下载原文件查看') from None

    @staticmethod
    def _cell(value):
        if value is None or isinstance(value,(str,bool,int)):return value[:1000] if isinstance(value,str) else value
        if isinstance(value,float):return value if math.isfinite(value) else str(value)
        if hasattr(value,'isoformat'):return value.isoformat()
        return str(value)[:1000]

    def _save(self,payload,receipt):
        self._db.execute('INSERT INTO receipts VALUES(?,?,?,?) ON CONFLICT(id) DO UPDATE SET receipt=excluded.receipt',
                         (receipt['id'],receipt['request_id'],_json(payload),_json(receipt)))
        self._db.commit()

    @staticmethod
    def _move(source,destination):
        if destination.exists():raise FileConflict('目标已存在，未覆盖原文件')
        # Windows rename is atomic and refuses existing destinations. On POSIX,
        # exclusive hard-link/file or copytree/directory avoids rename-overwrite.
        try:
            if os.name=='nt':source.rename(destination);return
            if source.is_file():os.link(source,destination);source.unlink();return
        except OSError as error:
            if destination.exists():raise FileConflict('目标已存在，未覆盖原文件') from None
            if error.errno not in (18,):raise
        before=_fingerprint(source)
        if source.is_dir():shutil.copytree(source,destination,copy_function=shutil.copy2)
        else:
            with source.open('rb') as src,destination.open('xb') as dst:shutil.copyfileobj(src,dst)
            shutil.copystat(source,destination)
        if _fingerprint(source)!=before:raise FileConflict('原内容在操作期间改变，已保留原文件；请检查目标')
        if source.is_dir():shutil.rmtree(source)
        else:source.unlink()

    def operation(self,body):
        if not isinstance(body,dict) or set(body)-{'operation','root_id','path','target','destination_root_id','receipt_id','request_id'}:raise ValueError('文件操作参数无效')
        action=body.get('operation')
        if action not in ('mkdir','rename','copy','move','trash','restore'):raise ValueError('不支持的文件操作')
        request=body.get('request_id') or str(uuid.uuid4())
        if not isinstance(request,str) or not re.fullmatch(r'[a-zA-Z0-9_-]{8,100}',request):raise ValueError('请求标识无效')
        payload={k:v for k,v in body.items() if k!='request_id'}
        with self._lock:
            previous=self._db.execute('SELECT payload,receipt FROM receipts WHERE request_id=?',(request,)).fetchone()
            if previous:
                if previous[0]!=_json(payload):raise FileConflict('同一请求标识不能用于不同文件操作')
                receipt=json.loads(previous[1])
                if receipt['status']=='completed':return receipt
                raise FileConflict('上次操作未确认完成，请检查目录后使用新请求重试')
            if action=='restore':return self._restore(body,payload,request)
            root_id=body.get('root_id');relative=_relative(body.get('path',''))
            if not relative:raise ValueError('不能操作登记目录本身')
            source=self.resolve(root_id,relative,must_exist=action!='mkdir')
            destination_root=body.get('destination_root_id',root_id)
            target=_relative(body.get('target','')) if action in ('rename','copy','move') else relative
            if action in ('rename','copy','move') and not target:raise ValueError('请输入目标相对路径')
            destination=self.resolve(destination_root,target,must_exist=False) if action!='trash' else self.recycle/uuid.uuid4().hex/source.name
            if action=='mkdir':destination=source
            if destination.exists():raise FileConflict('目标已存在，未覆盖原文件')
            if action!='trash' and not destination.parent.is_dir():raise ValueError('目标的上级文件夹不存在')
            if action in ('copy','move','rename') and (destination==source or source in destination.parents):raise ValueError('目标不能位于原文件夹内')
            if action=='rename' and destination_root!=root_id:raise ValueError('跨目录请使用移动操作')
            fingerprint=_fingerprint(source) if action!='mkdir' else ''
            receipt={'id':'fileop_'+uuid.uuid4().hex,'request_id':request,'operation':action,'status':'pending',
                'root_id':root_id,'path':relative,'destination_root_id':destination_root,'target':target,
                'created_at':time.time(),'restore_available':action=='trash','fingerprint':fingerprint}
            if action=='trash':receipt['recycle_id']=str(destination.relative_to(self.recycle))
            self._save(payload,receipt)
            try:
                if action=='mkdir':destination.mkdir()
                elif action=='copy':
                    if source.is_dir():shutil.copytree(source,destination,copy_function=shutil.copy2)
                    else:
                        with source.open('rb') as src,destination.open('xb') as dst:shutil.copyfileobj(src,dst)
                        shutil.copystat(source,destination)
                else:
                    if action=='trash':destination.parent.mkdir()
                    self._move(source,destination)
                receipt['status']='completed';self._save(payload,receipt);return receipt
            except Exception:
                receipt['status']='failed';receipt['detail']='操作未确认完成，请检查原目录和目标目录'
                self._save(payload,receipt);raise

    def _restore(self,body,payload,request):
        original=self._db.execute('SELECT payload,receipt FROM receipts WHERE id=?',(body.get('receipt_id'),)).fetchone()
        if not original:raise ValueError('回收记录不存在')
        old=json.loads(original[1])
        if old['operation']!='trash' or old['status']!='completed' or not old.get('restore_available'):raise FileConflict('此记录不能再次恢复')
        source=_plain(self.recycle,self.recycle/_relative(old['recycle_id']))
        destination=self.resolve(old['root_id'],old['path'],must_exist=False)
        if destination.exists():raise FileConflict('原地址已有文件，未覆盖现有内容')
        if not destination.parent.is_dir():raise ValueError('原上级文件夹不存在，请先恢复目录')
        if _fingerprint(source)!=old['fingerprint']:raise FileConflict('回收内容已改变，未自动恢复')
        receipt={'id':'fileop_'+uuid.uuid4().hex,'request_id':request,'operation':'restore','status':'pending',
                 'root_id':old['root_id'],'path':old['path'],'created_at':time.time(),'restore_available':False,'restored_receipt_id':old['id']}
        self._save(payload,receipt)
        try:
            self._move(source,destination)
            receipt['status']='completed';old['restore_available']=False
            self._save(json.loads(original[0]),old);self._save(payload,receipt);return receipt
        except Exception:
            receipt['status']='failed';self._save(payload,receipt);raise


class _HTMLText(HTMLParser):
    def __init__(self):super().__init__();self.parts=[];self.hidden=0
    def handle_starttag(self,tag,attrs):
        if tag in ('script','style'):self.hidden+=1
    def handle_endtag(self,tag):
        if tag in ('script','style') and self.hidden:self.hidden-=1
    def handle_data(self,data):
        if not self.hidden and data.strip():self.parts.append(data.strip())
