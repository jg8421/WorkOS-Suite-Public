"""Safe, portable deliverables. Text never becomes executable markup."""
import html
import io
from datetime import datetime
from pathlib import Path

ROOT=Path(__file__).resolve().parents[1]

def markdown(record):
 return '# '+record['title']+'\n\n'+record.get('body','')+'\n\n---\nLocal WorkOS · '+datetime.now().strftime('%Y-%m-%d')+'\n'

def html_report(record):
 title=html.escape(record['title'])
 lines=record.get('body','').splitlines()
 content=[]
 for i,line in enumerate(lines):
  escaped=html.escape(line)
  if line.startswith('## '):content.append(f'<h2 id="s{i}">{html.escape(line[3:])}</h2>')
  elif line.startswith('# '):content.append(f'<h2 id="s{i}">{html.escape(line[2:])}</h2>')
  elif line.strip():content.append(f'<p id="p{i}">{escaped}</p>')
  else:content.append('<div class="gap"></div>')
 report=f'''<!doctype html><html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1"><title>{title}</title><style>
*{{box-sizing:border-box}}body{{margin:0;background:#f4f5f7;color:#243345;font:16px/1.65 Arial,"Microsoft YaHei",sans-serif}}main{{max-width:1000px;margin:40px auto;padding:42px 48px;background:white;border:1px solid #dbe1e7;border-radius:12px}}h1{{font-size:30px;color:#153452;margin:0 0 18px}}h2{{font-size:22px;color:#153452;margin:28px 0 10px}}p{{white-space:pre-wrap;overflow-wrap:anywhere;margin:8px 0}}.notice{{padding:14px 18px;background:#fff8df;border:1px solid #d9c994;margin:18px 0}}.gap{{height:8px}}footer{{border-top:1px solid #dde3e9;margin-top:30px;padding-top:16px;color:#667687;font-size:13px}}@media(max-width:700px){{main{{margin:0;padding:70px 22px 80px;border:0;border-radius:0}}h1{{font-size:25px}}}}@media print{{body{{background:white}}main{{margin:0;border:0;padding:0;max-width:none}}p{{break-inside:avoid}}}}
</style></head><body><main id="report"><h1>{title}</h1><div class="notice" id="instructions">工作草稿，结论与数字请复核。默认阅读模式；点击“编辑正文”可改正文并添加批注。“保存 HTML 副本”下载包含修改和批注的独立文件；不回写平台数据库。</div><section id="body">{''.join(content)}</section><footer id="footer">Local WorkOS · {datetime.now().strftime('%Y-%m-%d')} · 请在外发前确认资料范围及敏感信息。</footer></main></body></html>'''
 editor=(ROOT/'web'/'report-editor.js').read_text(encoding='utf-8')
 # Prevent a script terminator even if the independently authored asset changes.
 editor=editor.replace('</', '<\\/')
 editor_css='''<style id="report-editor-style">
 body{padding-right:340px}#report-editor-ui{position:fixed;right:0;top:0;bottom:0;width:320px;overflow:auto;background:#f8fafc;border-left:1px solid #ccd5df;padding:18px;font:14px/1.5 Arial,sans-serif}#report-editor-ui strong,#report-editor-ui label{display:block;margin-bottom:10px}#report-editor-ui button{padding:8px 10px;margin:4px 5px 8px 0;cursor:pointer}#report-editor-ui textarea{width:100%;resize:vertical}#report-selected-quote,.report-note blockquote{white-space:pre-wrap;overflow-wrap:anywhere;background:#eef2f7;margin:8px 0;padding:8px}.report-note{border-top:1px solid #ccd5df;padding:10px 0}.report-note p{white-space:pre-wrap;overflow-wrap:anywhere}[contenteditable]{outline:1px dashed #8ba2bc;white-space:pre-wrap}h2{white-space:pre-wrap}@media(max-width:900px){body{padding-right:0}#report-editor-ui{position:static;width:auto;border-left:0;border-top:1px solid #ccd5df}}@media print{body{padding-right:0}#report-editor-ui{display:none}}
 </style>'''
 state='<script type="application/json" id="report-notes-data">{"version":1,"notes":[]}</script>'
 return report.replace('</head>',editor_css+'</head>').replace('</body>',state+'<script id="report-editor-script">'+editor+'</script></body>')


def normalize_minute_numbers(text):
 """Format explicit amounts without reinterpreting years or identifiers."""
 import re
 value=str(text or '')
 value=re.sub(r'(?<=\d)[ \t]*[%％]', '%', value)
 value=re.sub(r'(?<=[\d%])[ \t]*[－—–﹣][ \t]*(?=\d)', '-', value)
 value=re.sub(r'(?<=\d)[ \t]+(?=(?:μm|nm|mm|cm|km)(?![A-Za-z]))', '', value)
 def group(match):
  digits=match.group(1)
  return digits if digits.startswith('0') else format(int(digits), ',')
 value=re.sub(r'(?<![A-Za-z0-9_.,])([0-9]{4,})(?![0-9.,])(?=[ \t]*(?:百万元|万元|亿元|元|片|台|吨)(?:[^A-Za-z]|$))', group, value)
 return value

def normalize_minute_lines(summary):
 lines=[]
 for raw in str(summary or '').replace('\r\n','\n').replace('\r','\n').split('\n'):
  line=raw.rstrip()
  stripped=line.strip()
  if stripped:
   indent=line[:len(line)-len(line.lstrip())]
   line=indent+normalize_minute_numbers(stripped)
  lines.append(line)
 return '\n'.join(lines)

def expert_minutes_docx(title,summary,participants='',date_text='',experts=None,matrix=None,contents=None):
 from docx import Document
 from docx.shared import Pt,Cm
 from docx.enum.text import WD_TAB_ALIGNMENT,WD_TAB_LEADER
 from docx.oxml.ns import qn
 from io import BytesIO
 doc=Document();section=doc.sections[0];section.page_width=Cm(21);section.page_height=Cm(29.7);section.top_margin=Cm(2.2);section.bottom_margin=Cm(2.2);section.left_margin=Cm(2);section.right_margin=Cm(2)
 normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(10);normal.element.rPr.rFonts.set(qn('w:eastAsia'),'KaiTi')
 summary=normalize_minute_lines(summary);title=normalize_minute_numbers(title)
 p=doc.add_paragraph();run=p.add_run(title);run.bold=True;run.font.name='Arial';run.font.size=Pt(14);run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'KaiTi')
 if date_text:doc.add_paragraph(date_text)
 expert_list=experts if isinstance(experts,list) else []
 if expert_list:
  matrix=matrix if isinstance(matrix,dict) else {};topics=matrix.get('topics',[]);columns=matrix.get('experts',[]);cells=matrix.get('cells',[])
  if topics and columns and len(cells)==len(topics):
   table=doc.add_table(rows=1,cols=2);table.style='Table Grid';table.autofit=False;table.columns[0].width=Cm(3.3);table.columns[1].width=Cm(13.2)
   table.cell(0,0).width=Cm(3.3);table.cell(0,1).width=Cm(13.2);table.cell(0,0).text='议题';table.cell(0,1).text='专家口径'
   for topic,row in zip(topics,cells):
    for c,index in enumerate(columns):
     item=expert_list[index] if isinstance(index,int) and 0<=index<len(expert_list) else expert_list[min(c,len(expert_list)-1)]
     label='-'.join(str(item.get(k,'')).strip() for k in ('institution','title') if item.get(k)) or ('专家 '+str(c+1))
     values=row if isinstance(row,list) else []
     celltext=str(values[c]) if c<len(values) else '未提及'
     cells_row=table.add_row().cells;cells_row[0].width=Cm(3.3);cells_row[1].width=Cm(13.2);cells_row[0].text=str(topic);cells_row[1].text=label+'：'+celltext
   doc.add_page_break()
  if len(expert_list)>=4:
   from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
   from docx.oxml import OxmlElement
   doc.add_heading('目录',1);entries=contents if isinstance(contents,list) else []
   update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');doc.settings.element.append(update)
   for i,item in enumerate(expert_list):
    fallback='-'.join(str(item.get(k,'')).strip() for k in ('institution','title') if item.get(k)) or ('专家 '+str(i+1))
    toc=doc.add_paragraph();toc.paragraph_format.tab_stops.add_tab_stop(Cm(16.0), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    toc.add_run(str(entries[i] if i<len(entries) else fallback));toc.add_run('\t')
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),' PAGEREF expert_'+str(i)+' '+chr(92)+'h ');field.set(qn('w:dirty'),'true')
    run=OxmlElement('w:r');text=OxmlElement('w:t');text.text='更新域';run.append(text);field.append(run);toc._p.append(field)
   doc.add_page_break()
  for i,item in enumerate(expert_list):
   name='-'.join(str(item.get(k,'')).strip() for k in ('institution','title') if item.get(k)) or ('专家 '+str(i+1));heading=doc.add_heading(name,1)
   if len(expert_list)>=4:
    start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),str(i));start.set(qn('w:name'),'expert_'+str(i));end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(i));heading._p.insert(0,start);heading._p.append(end)
   date_value=str(item.get('date') or date_text or '').strip()
   if date_value:doc.add_paragraph(date_value)
   doc.add_heading('专家背景',2);doc.add_paragraph(str(item.get('background') or '未提及'));doc.add_heading('专家点评',2)
   for comment in item.get('comments',[]) if isinstance(item.get('comments'),list) else []:doc.add_paragraph(str(comment))
   doc.add_heading('访谈内容',2)
   content=str(item.get('content') or '').replace('\\n','\n').replace('\r\n','\n').replace('\r','\n')
   for line in content.splitlines():
    line=line.strip()
    if line:doc.add_paragraph(line if not line.startswith('- ') else '• '+line[2:])
   if i<len(expert_list)-1:doc.add_page_break()
  out=BytesIO();doc.save(out);return out.getvalue()
 lines=summary.replace('\r\n','\n').replace('\r','\n').split('\n')
 if participants:lines.insert(0,'专家身份：'+participants)
 for raw in lines:
  line=raw.strip()
  if not line:continue
  if line.startswith('【') and line.endswith('】'):doc.add_heading(line.strip('【】'),1)
  elif line.startswith('# '):doc.add_heading(line[2:].strip(),1)
  elif line.startswith('## '):doc.add_heading(line[3:].strip(),2)
  elif line.startswith(('➢','o ','• ','- ')):
   p=doc.add_paragraph(style='Normal');p.paragraph_format.left_indent=Cm(1.0 if line.startswith('➢') else .65 if line.startswith('o ') else .35);p.paragraph_format.first_line_indent=Cm(-.3);p.add_run(('• '+line[2:] if line.startswith('- ') else line))
  else:doc.add_paragraph(line)
 out=BytesIO();doc.save(out);return out.getvalue()
def valuation_xlsx(method, assumptions, result):
    if method == 'investor_return':
        from .return_excel import workbook
        return workbook(assumptions)[1]
    import io, json
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    wb = Workbook(); summary = wb.active; summary.title = "Summary"
    from openpyxl.workbook.properties import CalcProperties
    wb.calculation=CalcProperties(calcMode="auto",fullCalcOnLoad=True,forceFullCalc=True)
    summary.append(["Local WorkOS Valuation Model"]); summary.append(["Method", result.get("method_label", method)]); summary.append(["Currency / Unit", str(assumptions.get("currency", ""))+" / "+str(assumptions.get("unit", ""))]); summary.append(["Formula", result.get("formula", "")]); summary.append([]); summary.append(["Metric", "Value / Formula"])
    ass = wb.create_sheet("Assumptions"); ass.append(["Assumption", "Input"])
    for key, value in assumptions.items():
        if key != "forecasts":
            if key in ("entry_date","exit_date","valuation_date") and isinstance(value,str):
                from datetime import date
                value=date.fromisoformat(value)
            ass.append([key, json.dumps(value, ensure_ascii=False) if isinstance(value, (dict,list)) else value])
    ass.freeze_panes="A2"; ass.column_dimensions["A"].width=30; ass.column_dimensions["B"].width=28
    def ref(key):
        row = next((r for r in range(2, ass.max_row+1) if ass.cell(r,1).value == key), None)
        return "Assumptions!$B$" + str(row) if row else "0"
    if method in ("net_income", "ps"):
        input_key = "net_income" if method == "net_income" else "revenue"; multiple_key = "pe_multiple" if method == "net_income" else "ps_multiple"
        summary.append(["Equity Value", "="+ref(input_key)+"*"+ref(multiple_key)])
        if assumptions.get("diluted_shares"): summary.append(["Implied Value / Share", "=B7/"+ref("diluted_shares")])
    elif method == "dcf":
        forecasts = assumptions.get("forecasts") or []; ws = wb.create_sheet("DCF_Forecast")
        ws.append(["Year","EBIT","Tax Rate","Cash Tax","NOPAT","D&A","CapEx","ΔNWC","FCFF","Discount Period","Discount Factor","PV FCFF"])
        for idx, row in enumerate(forecasts, 2):
            ebit = row.get("ebit") if row.get("ebit") is not None else row.get("revenue",0)*row.get("ebit_margin",0)
            tax = row.get("tax_rate") if row.get("tax_rate") is not None else "="+ref("tax_rate")
            period = '=IF(%s="mid_year",%s,%s)'%(ref("discount_timing"),idx-1.5,idx-1)
            ws.append([row.get("year"), ebit, tax, "=MAX(0,B%d)*C%d"%(idx,idx), "=B%d-D%d"%(idx,idx), row.get("da"), row.get("capex"), row.get("delta_nwc"), "=E%d+F%d-G%d-H%d"%(idx,idx,idx,idx), period, "=1/(1+%s)^J%d"%(ref("wacc"),idx), "=I%d*K%d"%(idx,idx)])
        if forecasts:
            n = len(forecasts)+1; last = n; tvrow = n+3
            if assumptions.get("terminal_method") == "perpetuity": terminal = "=I%d*(1+%s)/(%s-%s)"%(last,ref("terminal_growth"),ref("wacc"),ref("terminal_growth"))
            else: terminal = "=(B%d+F%d)*%s"%(last,last,ref("terminal_multiple"))
            ws.cell(tvrow,1,"Terminal Value"); ws.cell(tvrow,2,terminal); ws.cell(tvrow+1,1,"PV Terminal"); ws.cell(tvrow+1,2,"=B%d/(1+%s)^J%d"%(tvrow,ref("wacc"),last))
            evrow = summary.max_row+1; summary.append(["Enterprise Value", "=SUM(DCF_Forecast!L2:L%d)+DCF_Forecast!B%d"%(n,tvrow+1)]); summary.append(["Equity Value", "=B%d-%s-%s"%(evrow,ref("net_debt"),ref("minority_interest"))])
        ws.freeze_panes="A2"
        for col, width in enumerate((16,16,14,16,16,14,14,14,16,18,18,18),1): ws.column_dimensions[get_column_letter(col)].width=width
    elif method == "lbo":
        forecasts = assumptions.get("forecasts") or []; ws = wb.create_sheet("LBO_Model")
        ws.append(["Period","EBITDA","D&A","CapEx","ΔNWC","Tax Rate","Interest Rate","Mandatory Amort.","Cash Sweep %","Opening Debt","Opening Cash","Interest","Tax","Cash Before Debt","Mandatory Due","Mandatory Paid","Sweep","Ending Debt","Ending Cash","Funding Gap","Mandatory Shortfall","Cash Floor Shortfall"])
        for idx, row in enumerate(forecasts,2):
            vals = [row.get("year"),row.get("ebitda"),row.get("da"),row.get("capex"),row.get("delta_nwc")]+[row[key] if row.get(key) is not None else "="+ref(key) for key in ("tax_rate","interest_rate","mandatory_amortization","cash_sweep_pct")]
            f = ["="+ref("entry_debt") if idx==2 else "=R%d"%(idx-1), "="+ref("initial_cash") if idx==2 else "=S%d"%(idx-1), "=J%d*G%d"%(idx,idx), "=MAX(0,B%d-C%d-L%d)*F%d"%(idx,idx,idx,idx), "=B%d-M%d-L%d-D%d-E%d"%(idx,idx,idx,idx,idx), "=MIN(J%d,H%d)"%(idx,idx), "=MIN(O%d,MAX(0,K%d+N%d-%s))"%(idx,idx,idx,ref("minimum_cash")), "=MIN(MAX(0,J%d-P%d),MAX(0,K%d+N%d-P%d-%s)*I%d)"%(idx,idx,idx,idx,idx,ref("minimum_cash"),idx), "=MAX(0,J%d-P%d-Q%d)"%(idx,idx,idx), "=MAX(0,K%d+N%d-P%d-Q%d)"%(idx,idx,idx,idx), "=U%d+V%d"%(idx,idx), "=O%d-P%d"%(idx,idx), "=MAX(0,%s-(K%d+N%d-P%d))"%(ref("minimum_cash"),idx,idx,idx)]
            ws.append(vals+f)
        ws.freeze_panes="A2"
        if forecasts:
            n = len(forecasts)+1
            for col,label in enumerate(("Opening Debt","Opening Cash","Interest","Tax","Cash Before Debt","Mandatory Due","Mandatory Paid","Sweep","Ending Debt","Ending Cash","Funding Gap"),10): ws.cell(1,col,label)
            er=n+2; ws.cell(er,1,"Formula Exit EV"); ws.cell(er,2,"=B%d*%s"%(n,ref("exit_multiple")))
            tr=n+12; tor=n+13; owr=n+14; sellerpr=n+15; sellerown=n+16
            pr=n+3; ws.cell(pr,1,"Formula Sponsor Proceeds"); ws.cell(pr,2,"=B%d*B%d"%(tr,owr))
            sr=n+4; ws.cell(sr,1,"Formula Sponsor Equity"); ws.cell(sr,2,"="+ref("entry_ev")+"+"+ref("entry_fees")+"+"+ref("initial_cash")+"-"+ref("entry_debt")+"-"+ref("seller_rollover"))
            mr=n+5; ws.cell(mr,1,"Formula MOIC"); ws.cell(mr,2,"=B%d/B%d"%(pr,sr))
            from datetime import date as _date
            ws.cell(n+6,1,"Entry Date"); ws.cell(n+6,2,"="+ref("entry_date"));ws.cell(n+6,2).number_format="yyyy-mm-dd"
            ws.cell(n+7,1,"Exit Date"); ws.cell(n+7,2,"="+ref("exit_date"));ws.cell(n+7,2).number_format="yyyy-mm-dd"
            ws.cell(n+8,1,"Formula IRR"); ws.cell(n+8,2,"=B%d^(365/(B%d-B%d))-1"%(mr,n+7,n+6))
            ws.cell(n+10,1,"Python Ground Truth — Sponsor Proceeds"); ws.cell(n+10,2,result.get("sponsor_proceeds")); ws.cell(n+11,1,"Python Ground Truth — IRR"); ws.cell(n+11,2,result.get("irr"))
            ws.cell(tr,1,"Formula Total Exit Proceeds"); ws.cell(tr,2,"=MAX(0,B%d-R%d+S%d-%s)"%(er,n,n,ref("exit_fees")))
            ws.cell(tor,1,"Formula Total Entry Equity"); ws.cell(tor,2,"=B%d+%s"%(sr,ref("seller_rollover")))
            ws.cell(owr,1,"Formula Sponsor Ownership"); ws.cell(owr,2,"=B%d/B%d"%(sr,tor));ws.cell(owr,2).number_format="0.0%"
            ws.cell(sellerpr,1,"Formula Seller Proceeds");ws.cell(sellerpr,2,"=B%d-B%d"%(tr,pr))
            ws.cell(sellerown,1,"Formula Seller Ownership");ws.cell(sellerown,2,"=1-B%d"%owr);ws.cell(sellerown,2).number_format="0.0%"
            ws.cell(n+18,1,"Ownership / financing assumptions");ws.cell(n+18,2,result.get("warning", ""))
        for col in range(1,23): ws.column_dimensions[get_column_letter(col)].width=16
    rows = result.get("forecast") or result.get("rows")
    if isinstance(rows, list):
        output=wb.create_sheet("Calculated_Output"); keys=sorted({k for row in rows if isinstance(row,dict) for k in row}); output.append(keys)
        for row in rows:
            if isinstance(row,dict): output.append([row.get(k) for k in keys])
    if method=="lbo" and forecasts:
        for label,rownum in (("Sponsor Proceeds",pr),("Sponsor Equity",sr),("MOIC",mr),("IRR",n+8),
                             ("Total Exit Proceeds",tr),("Seller Proceeds",sellerpr),("Sponsor Ownership",owr),("Seller Ownership",sellerown)):
            summary.append([label,"=LBO_Model!B"+str(rownum)])
    if method=="ps" and "net_debt" in assumptions:summary.append(["Enterprise Value","=B7+"+ref("net_debt")])
    summary.cell(6,3,"Python Snapshot — original inputs")
    mapping={"Equity Value":"equity_value","Enterprise Value":"enterprise_value" if method=="dcf" else "implied_enterprise_value","Implied Value / Share":"implied_value_per_share","Sponsor Proceeds":"sponsor_proceeds","Sponsor Equity":"entry_sponsor_equity","MOIC":"moic","IRR":"irr","Total Exit Proceeds":"total_exit_proceeds","Seller Proceeds":"seller_proceeds","Sponsor Ownership":"sponsor_ownership","Seller Ownership":"seller_ownership"}
    for rownum in range(7,summary.max_row+1):
        key=mapping.get(summary.cell(rownum,1).value)
        if key in result:summary.cell(rownum,3,result[key])
    summary.column_dimensions["A"].width=30;summary.column_dimensions["B"].width=24;summary.column_dimensions["C"].width=32
    for sheet in wb.worksheets:
        sheet.sheet_view.showGridLines=False
        if sheet.max_row: sheet.freeze_panes=sheet.freeze_panes or "A2"
        for cell in sheet[1]: cell.font=Font(name="Arial",bold=True,color="FFFFFF"); cell.fill=PatternFill("solid",fgColor="17365D"); cell.alignment=Alignment(wrap_text=True)
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is not None:
                    cell.font=Font(name="Arial",size=10,bold=cell.row==1,color="FFFFFF" if cell.row==1 else "243345")
                    cell.alignment=Alignment(vertical="top",wrap_text=True)
    out=io.BytesIO();wb.save(out);return out.getvalue()


def _slide_pages(sections, maximum=200):
 """Paginate every paragraph; reject excessive decks instead of truncating."""
 import unicodedata
 pages=[]
 for heading,items in sections:
  page=[];weight=0;part=1
  for item in items:
   chunks=[];chunk="";width=0
   for char in item:
    cost=2 if unicodedata.east_asian_width(char) in ("W","F") else 1
    if width+cost>100 and chunk:
     chunks.append(chunk);chunk="";width=0
    chunk+=char;width+=cost
   if chunk:chunks.append(chunk)
   for chunk in chunks:
    if weight>=8:
     pages.append((heading if part==1 else heading+"（续"+str(part)+"）",page))
     if len(pages)>maximum:raise ValueError("正文需超过200页；请拆分交付物。未生成截断的PPT。")
     page=[];weight=0;part+=1
    page.append(chunk);weight+=1
  if page or not items:pages.append((heading if part==1 else heading+"（续"+str(part)+"）",page))
  if len(pages)>maximum:raise ValueError("正文需超过200页；请拆分交付物。未生成截断的PPT。")
 return pages

def pptx_report(record):
 from pptx import Presentation
 from pptx.util import Inches,Pt
 from pptx.dml.color import RGBColor
 from pptx.enum.text import PP_ALIGN
 title=str(record.get('title') or 'Research Material')[:200]
 body=str(record.get('body') or '').replace('\r\n','\n').replace('\r','\n')
 if '\n' not in body:body=body.replace('\\n','\n')
 prs=Presentation();prs.slide_width= Inches(13.333);prs.slide_height= Inches(7.5)
 navy=RGBColor(23,54,93);blue=RGBColor(47,91,147);cream=RGBColor(250,246,235);muted=RGBColor(95,110,125)
 def add_title(slide,text,subtitle=None):
  box=slide.shapes.add_textbox(Inches(.65),Inches(.35),Inches(12),Inches(.8));p=box.text_frame.paragraphs[0];p.text=text;p.font.name='Arial';p.font.size=Pt(25);p.font.bold=True;p.font.color.rgb=navy
  if subtitle:
   b=slide.shapes.add_textbox(Inches(.68),Inches(1.12),Inches(12),Inches(.35));q=b.text_frame.paragraphs[0];q.text=subtitle;q.font.name='Arial';q.font.size=Pt(10);q.font.color.rgb=muted
 def add_footer(slide,n):
  bar=slide.shapes.add_shape(1,Inches(.65),Inches(7.12),Inches(12),Inches(.02));bar.fill.solid();bar.fill.fore_color.rgb=navy;bar.line.fill.background()
  box=slide.shapes.add_textbox(Inches(.68),Inches(7.16),Inches(12),Inches(.2));p=box.text_frame.paragraphs[0];p.text='Local WorkOS · 工作草稿 · '+str(n);p.font.name='Arial';p.font.size=Pt(8);p.font.color.rgb=muted
 # Cover
 slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.background.fill.solid();slide.background.fill.fore_color.rgb=cream
 band=slide.shapes.add_shape(1,Inches(0),Inches(0),Inches(.18),Inches(7.5));band.fill.solid();band.fill.fore_color.rgb=navy;band.line.fill.background()
 box=slide.shapes.add_textbox(Inches(.9),Inches(1.9),Inches(11.5),Inches(1.3));p=box.text_frame.paragraphs[0];p.text=title;p.font.name='Arial';p.font.size=Pt(30);p.font.bold=True;p.font.color.rgb=navy
 sub=slide.shapes.add_textbox(Inches(.95),Inches(3.3),Inches(11),Inches(.7));p=sub.text_frame.paragraphs[0];p.text='讨论材料 / 投资研究';p.font.name='Arial';p.font.size=Pt(18);p.font.color.rgb=blue
 note=slide.shapes.add_textbox(Inches(.95),Inches(6.65),Inches(11),Inches(.35));p=note.text_frame.paragraphs[0];p.text='请核对数据、来源、保密范围和结论口径后再外发。';p.font.name='Arial';p.font.size=Pt(10);p.font.color.rgb=muted
 from .deck_blocks import parse_blocks
 pages=[]
 for block in parse_blocks(body):
  head=block["heading"]
  if block["kind"]=="text":
   pages.extend(("text",title,items) for title,items in _slide_pages([(head,block["items"])]))
  elif block["kind"]=="table":
   rows=block["rows"]
   for start in range(0,max(1,len(rows)),6):
    pages.append(("table",head if not start else head+"（续"+str(start//6+1)+"）",{"header":block["header"],"rows":rows[start:start+6]}))
  else:pages.append(("chart",block["data"].get("title") or head,block["data"]))
  if len(pages)>200:raise ValueError("正文需超过200页，请拆分；未生成截断PPT")
 for index,(kind,head,items) in enumerate(pages,1):
  slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.background.fill.solid();slide.background.fill.fore_color.rgb=RGBColor(255,255,255);add_title(slide,head,'证据与结论需回到来源材料复核')
  if kind=="chart":
   from pptx.chart.data import CategoryChartData
   from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
   data=CategoryChartData();data.categories=items["categories"]
   for series in items["series"]:data.add_series(series["name"],series["values"])
   chart_type={"column":XL_CHART_TYPE.COLUMN_CLUSTERED,"bar":XL_CHART_TYPE.BAR_CLUSTERED,"line":XL_CHART_TYPE.LINE_MARKERS}[items["type"]]
   chart=slide.shapes.add_chart(chart_type,Inches(.9),Inches(1.65),Inches(11.6),Inches(4.9),data).chart
   chart.has_legend=len(items["series"])>1
   if chart.has_legend:chart.legend.position=XL_LEGEND_POSITION.BOTTOM
   source=items.get("source") or "来源未提供；数据来自交付物中的明确图表数据块"
   source_box=slide.shapes.add_textbox(Inches(.9),Inches(6.64),Inches(11.6),Inches(.35))
   source_box.text_frame.paragraphs[0].text=source;source_box.text_frame.paragraphs[0].font.size=Pt(10)
   slide.notes_slide.notes_text_frame.text="图表来源："+source
   add_footer(slide,index+1);continue
  if kind=="table":
   rows=[items["header"]]+items["rows"];table=slide.shapes.add_table(len(rows),len(items["header"]),Inches(.9),Inches(1.65),Inches(11.6),Inches(4.9)).table
   for ri,row in enumerate(rows):
    for ci,value in enumerate(row):
     cell=table.cell(ri,ci);cell.text=value;cell.text_frame.word_wrap=True
     cell.fill.solid();cell.fill.fore_color.rgb=navy if ri==0 else cream
     for paragraph in cell.text_frame.paragraphs:
      paragraph.font.name="Arial";paragraph.font.size=Pt(12);paragraph.font.bold=ri==0;paragraph.font.color.rgb=RGBColor(255,255,255) if ri==0 else RGBColor(36,51,69)
   slide.notes_slide.notes_text_frame.text="可编辑表格；数值及口径保持原交付正文，请核对来源。"
   add_footer(slide,index+1);continue
  box=slide.shapes.add_textbox(Inches(.9),Inches(1.55),Inches(11.6),Inches(5.25));tf=box.text_frame;tf.clear();tf.word_wrap=True
  for j,line in enumerate(items):
    p=tf.paragraphs[0] if j==0 else tf.add_paragraph();text=line;level=0;bullet=False
    if text.startswith('➢'):level=2;bullet=True;text=text[1:].strip()
    elif text.startswith('o '):level=1;bullet=True;text=text[2:].strip()
    elif text.startswith(('• ','- ')):bullet=True;text=text[2:].strip()
    p.text=text;p.level=level;p.font.name='Arial';p.font.size=Pt(16 if level==0 else 14);p.font.color.rgb=RGBColor(36,51,69);p.space_after=Pt(10);p.text=('• '+text) if bullet and not level else text

  add_footer(slide,index+1)
 out=io.BytesIO();prs.save(out);return out.getvalue()

def docx_report(record):
 try:
  from docx import Document
  from docx.shared import Pt,Cm,RGBColor
  from docx.oxml import OxmlElement
  from docx.oxml.ns import qn
 except ImportError:
  # The HTTP adapter identifies optional-library setup failures by library name.
  # Keep the original ImportError rather than turning it into a task-input error.
  raise
 doc=Document()
 section=doc.sections[0];section.top_margin=Cm(1.8);section.bottom_margin=Cm(1.8);section.left_margin=Cm(2);section.right_margin=Cm(2)
 normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(10.5)
 normal.element.rPr.rFonts.set(qn('w:eastAsia'),'KaiTi')
 doc.add_heading(record['title'],0)
 for line in record.get('body','').splitlines():
  if line.startswith('#'):doc.add_heading(line.lstrip('# ').strip(),min(len(line)-len(line.lstrip('#')),3))
  else:doc.add_paragraph(line)
 doc.add_paragraph('工作草稿，请复核来源、数字和外发范围。')
 out=io.BytesIO();doc.save(out);return out.getvalue()
