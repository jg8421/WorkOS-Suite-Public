"""Portable XLSX authoring; native Excel remains the calculation authority.

No dependency on a Codex installation, Node or proprietary artifact runtime.
The authored workbook is recalculated and independently verified by return_excel.
"""
from pathlib import Path


def author(assumptions, result, destination):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.workbook.properties import CalcProperties

    a, r = assumptions, result
    wb = Workbook()
    wb.calculation = CalcProperties(calcId=0, fullCalcOnLoad=True, forceFullCalc=True)
    sheet = wb.active
    sheet.title = 'Returns'
    sheet.sheet_view.showGridLines = False
    def text(value):
        value = str(value if value is not None else '')
        return "'" + value if value[:1] in ('=', '+', '@', '-') else value
    def put(cell, value):
        sheet[cell] = text(value) if isinstance(value, str) else value
    def formula(cell, value):
        sheet[cell] = value
        sheet[cell].font = Font(name='Arial', size=10, color='000000')
    def date(cell, value):
        year, month, day = map(int, value.split('-'))
        formula(cell, f'=DATE({year},{month},{day})')
        sheet[cell].number_format = 'yyyy-mm-dd'

    labels = {'A1': 'Investor Returns', 'A2': f'{r["currency"]} · {r["unit"]} · gross investor cash flows',
              'A4': 'MOC / MOIC', 'A5': 'IRR (XIRR)', 'A6': 'Total invested', 'A7': 'Total received',
              'A9': 'Inputs — blue cells are editable', 'D10': 'Entry ownership',
              'D11': 'Dilution events (excluding IPO input)', 'D12': 'Exit ownership',
              'D13': 'Exit company equity', 'D14': 'Investor exit proceeds',
              'A24': 'Date', 'B24': 'Investment', 'C24': 'Distribution',
              'D24': 'Net cash flow', 'E24': 'Event'}
    for cell, value in labels.items(): put(cell, value)
    inputs = [('Entry equity value', a.get('entry_equity_value')), ('Initial investment', a.get('investment_amount')),
              ('Entry ownership (if supplied)', a.get('entry_ownership')), ('IPO dilution', a.get('ipo_dilution', 0)),
              ('Final ownership (if supplied)', a.get('exit_ownership')), ('Exit net income', a.get('exit_net_income')),
              ('Exit P/E', a.get('exit_pe_multiple')), ('Exit equity value (if supplied)', a.get('exit_equity_value')),
              ('Exit proceeds (if supplied)', a.get('exit_proceeds')), ('Entry valuation basis', a.get('entry_valuation_basis')),
              ('Entry date', r['entry_date']), ('Exit date', r['exit_date'])]
    for row, (label, value) in enumerate(inputs, 10):
        put(f'A{row}', label); put(f'B{row}', value)
    date('B20', r['entry_date']); date('B21', r['exit_date'])
    formula('E10', '=IF(ISNUMBER(B12),B12,IF(ISNUMBER(B10),B11/IF(B19="pre_money",B10+B11,B10),0))')
    for row, amount in enumerate(a.get('dilution_events') or [], 10): put(f'G{row}', amount)
    factors = ','.join(f'1-G{row}' for row in range(10, 30))
    formula('E12', f'=IF(ISNUMBER(B14),B14,E10*(1-B13)*PRODUCT({factors}))')
    formula('E13', '=IF(ISNUMBER(B17),B17,B15*B16)')
    formula('E14', '=IF(ISNUMBER(B18),B18,E13*E12)')
    full = a.get('cash_flows') is not None
    if full:
        for cell in ('E10', 'E12', 'E13', 'E14'): put(cell, 'Cash-flow schedule below')
    flows = r['cash_flows'] if full else [
        {'date': r['entry_date'], 'amount': -a['investment_amount'], 'label': 'Initial investment'},
        *(a.get('interim_cash_flows') or []),
        {'date': r['exit_date'], 'amount': r['exit_proceeds'], 'label': 'Exit proceeds'}]
    for index, flow in enumerate(flows):
        row = 25 + index
        date(f'A{row}', flow['date'])
        put(f'B{row}', max(0, -flow['amount'])); put(f'C{row}', max(0, flow['amount']))
        put(f'E{row}', flow.get('label', ''))
        if not full and index == 0:
            formula(f'A{row}', '=B20'); formula(f'B{row}', '=B11')
        if not full and index == len(flows) - 1:
            formula(f'A{row}', '=B21'); formula(f'C{row}', '=E14')
        formula(f'D{row}', f'=C{row}-B{row}')
    end = 24 + len(flows)
    formula('B6', f'=SUM(B25:B{end})'); formula('B7', f'=SUM(C25:C{end})')
    formula('B4', '=B7/B6')
    if r['irr'] is not None: formula('B5', f'=XIRR(D25:D{end},A25:A{end},{r["irr"]})')
    else: put('B5', 'No finite XIRR (no distributions)')
    put(f'A{end+3}', 'Scope'); put(f'B{end+3}', r['warning'])
    for row in sheet.iter_rows(min_row=1, max_row=end+4, max_col=7):
        for cell in row:
            cell.font = Font(name='Arial', size=10, color='000000')
            cell.alignment = Alignment(vertical='center')
    for row in range(1, end+5): sheet.row_dimensions[row].height = 22
    for start, stop, color, white in ((1, 2, '173B42', True), (4, 7, 'E7F4EF', False), (24, 24, 'DDE8EB', False)):
        for row in sheet.iter_rows(min_row=start, max_row=stop, max_col=7 if start == 1 else 2 if start == 4 else 5):
            for cell in row:
                cell.fill = PatternFill('solid', fgColor=color)
                cell.font = Font(name='Arial', size=10, color='FFFFFF' if white else '000000', bold=start == 4)
    money = '#,##0.00;(#,##0.00);"-"'
    for row in range(6, 20): sheet[f'B{row}'].number_format = money
    for row in range(10, 20): sheet[f'B{row}'].font = Font(name='Arial', size=10, color='0000FF')
    for row in range(10, 30):
        sheet[f'G{row}'].font = Font(name='Arial', size=10, color='0000FF')
        sheet[f'G{row}'].number_format = '0.0%'
    for cell in ('B12', 'B13', 'B14', 'E10', 'E12'): sheet[cell].number_format = '0.0%'
    for row in range(25, end+1):
        for column in ('B', 'C', 'D'): sheet[f'{column}{row}'].number_format = money
        for column in ('A', 'B', 'C'):
            cell = sheet[f'{column}{row}']
            cell.font = Font(name='Arial', size=10, color='000000' if cell.data_type == 'f' else '0000FF')
    for cell in ('B20', 'B21'): sheet[cell].font = Font(name='Arial', size=10, color='0000FF')
    sheet['B4'].number_format = '0.00"x"'; sheet['B5'].number_format = '0.0%'
    for col, width in {'A': 33, 'B': 25, 'C': 18, 'D': 34, 'E': 27, 'G': 15}.items(): sheet.column_dimensions[col].width = width
    sheet.freeze_panes = 'A8'
    sheet.print_options.horizontalCentered = True
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1; sheet.page_setup.fitToHeight = 0
    sheet.print_area = f'A1:G{end+4}'
    sources = wb.create_sheet('Sources'); sources.sheet_view.showGridLines = False
    sources.append(['Input / evidence', 'Source location'])
    for key, value in list((a.get('assumption_sources') or {}).items())[:80]:
        import json
        sources.append([text(key), text(json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value)])
    sources.column_dimensions['A'].width = 30; sources.column_dimensions['B'].width = 95
    for row in sources:
        for cell in row: cell.font = Font(name='Arial', size=10, color='008000')
    wb.save(Path(destination))

