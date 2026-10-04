"""Real loopback HTTP tests, ephemeral ports, no browser or external model calls."""
import base64
import html
import http.client
from http.server import ThreadingHTTPServer
import json
from pathlib import Path
from tempfile import TemporaryDirectory
import threading
import unittest
from unittest.mock import patch

from workos.exports import html_report, markdown
from workos.server import Application, Handler


class ServerTests(unittest.TestCase):
    def setUp(self):
        self.tmp = TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        # Prevent discovery of the real workstation memory root altogether.
        with patch('workos.server.find_root', return_value=None):
            self.app = Application(Path(self.tmp.name) / 'data', port=0)
            self.app.public_auth_mode='access'
        self.addCleanup(self.close_stores)
        self.httpd = ThreadingHTTPServer(('127.0.0.1', 0), Handler)
        self.httpd.daemon_threads = True
        self.app.port = self.httpd.server_address[1]
        self.httpd.app = self.app
        self.thread = threading.Thread(target=self.httpd.serve_forever, kwargs={'poll_interval': 0.02}, daemon=True)
        self.thread.start()
        self.addCleanup(self.stop_server)

    def close_stores(self):
        self.app.close()

    def stop_server(self):
        self.httpd.shutdown()
        self.httpd.server_close()
        self.thread.join(timeout=3)
        self.assertFalse(self.thread.is_alive(), 'HTTP test thread failed to stop')

    def request(self, method, path, body=None, workspace='personal', headers=None, csrf=True):
        connection = http.client.HTTPConnection('127.0.0.1', self.app.port, timeout=3)
        actual = {'X-Workspace': workspace}
        if csrf:
            actual['X-CSRF-Token'] = self.app.csrf
        if body is not None:
            actual['Content-Type'] = 'application/json'
            body = json.dumps(body, ensure_ascii=False).encode('utf-8')
        actual.update(headers or {})
        try:
            connection.request(method, path, body=body, headers=actual)
            response = connection.getresponse()
            raw = response.read()
            mime = response.getheader('Content-Type', '')
            parsed = json.loads(raw) if mime.startswith('application/json') else raw.decode('utf-8') if mime.startswith('text/') else raw
            return response.status, parsed
        finally:
            connection.close()

    def create(self, collection, body, workspace='personal'):
        status, result = self.request('POST', '/api/' + collection, body, workspace=workspace)
        self.assertEqual(status, 201, result)
        return result

    def test_bootstrap_and_crud(self):
        status, boot = self.request('GET', '/api/bootstrap', csrf=False)
        self.assertEqual(status, 200)
        self.assertEqual(boot['csrf'], self.app.csrf)
        self.assertEqual(boot['workspace'], 'personal')
        self.assertFalse(boot['memory_root_available'])
        project = self.create('projects', {'name': 'HTTP合成项目'})
        task = self.create('tasks', {'title': 'HTTP任务', 'project_id': project['id']})
        status, result = self.request('PATCH', '/api/tasks/' + task['id'], {'status': '完成'})
        self.assertEqual(status, 200, result)
        status, state = self.request('GET', '/api/state')
        self.assertEqual(state['tasks'][0]['status'], '完成')
        self.assertEqual(self.request('DELETE', '/api/projects/' + project['id'])[0], 400)
        self.assertEqual(self.request('DELETE', '/api/tasks/' + task['id'])[0], 200)
        self.assertEqual(self.request('DELETE', '/api/projects/' + project['id'])[0], 200)

    def test_csrf_required_for_each_mutation(self):
        project = self.create('projects', {'name': '保留'})
        for method, path, body in [('POST', '/api/projects', {'name': '拒绝'}),
                                   ('PATCH', '/api/projects/' + project['id'], {'name': '拒绝'}),
                                   ('DELETE', '/api/projects/' + project['id'], None),
                                   ('POST', '/api/shutdown', {})]:
            for token in (None, 'wrong-token'):
                with self.subTest(method=method, token=token):
                    # Security validation precedes parsing; omit an unread body to avoid Windows TCP close/reset races.
                    status, result = self.request(method, path, None, csrf=False,
                                                  headers={} if token is None else {'X-CSRF-Token': token})
                    self.assertEqual(status, 403)
                    self.assertIn('error', result)
        self.assertEqual(self.request('GET', '/api/state')[1]['projects'][0]['name'], '保留')

    def test_host_and_origin_validation(self):
        for headers in ({'Host': 'attacker.invalid'}, {'Host': '127.0.0.1:1'},
                        {'Origin': 'https://attacker.invalid'}, {'Origin': 'null'}):
            for method, body in [('GET', None), ('POST', {'name': '拒绝'})]:
                with self.subTest(headers=headers, method=method):
                    path = '/api/state' if method == 'GET' else '/api/projects'
                    self.assertEqual(self.request(method, path, None, headers=headers)[0], 403)
        good = {'Host': 'localhost:' + str(self.app.port), 'Origin': 'http://localhost:' + str(self.app.port)}
        self.assertEqual(self.request('GET', '/api/state', headers=good)[0], 200)

    def test_workspace_isolation_and_restore(self):
        project = self.create('projects', {'name': '个人合成'})
        personal = self.request('GET', '/api/backup')[1]
        demo_before = self.request('GET', '/api/backup', workspace='demo')[1]['data']
        self.assertNotIn(project['id'], {p['id'] for p in demo_before['projects']})
        self.assertEqual(self.request('POST', '/api/tasks', {'title': '跨区', 'project_id': project['id']}, workspace='demo')[0], 400)
        self.assertEqual(self.request('POST', '/api/restore', {'backup': personal, 'confirm': True}, workspace='demo')[0], 400)
        self.request('PATCH', '/api/projects/' + project['id'], {'name': '修改'})
        self.assertEqual(self.request('POST', '/api/restore', {'backup': personal, 'confirm': False})[0], 400)
        self.assertEqual(self.request('POST', '/api/restore', {'backup': personal, 'confirm': True})[0], 200)
        self.assertEqual(self.request('GET', '/api/backup')[1]['data'], personal['data'])
        self.assertEqual(self.request('GET', '/api/backup', workspace='demo')[1]['data'], demo_before)
        before = self.request('GET', '/api/backup')[1]['data']
        self.assertEqual(self.request('POST', '/api/restore', {'backup': {}, 'confirm': True})[0], 400)
        self.assertEqual(self.request('GET', '/api/backup')[1]['data'], before)
        self.assertEqual(self.request('GET', '/api/state', workspace='invalid')[0], 400)

    def test_original_upload_download_versions_and_text_edit(self):
        project = self.create('projects', {'name': '合成版本项目'})
        originals = []
        for version in (1, 2):
            raw = ('Synthetic IC memo version ' + str(version)).encode()
            status, doc = self.request('POST', '/api/upload', {'name': 'Memo_v' + str(version) + '.txt',
                'base64': base64.b64encode(raw).decode(), 'project_id': project['id'],
                'source_ref': 'Materials/Memo_v' + str(version) + '.txt'})
            self.assertEqual(status, 201, doc)
            self.assertEqual(self.request('GET', '/api/documents/' + doc['id'] + '/original'), (200, raw))
            self.assertEqual(self.request('GET', '/api/documents/' + doc['id'] + '/original', workspace='demo')[0], 404)
            originals.append((doc, raw))
        self.assertNotEqual(originals[0][0]['id'], originals[1][0]['id'])
        self.assertEqual(originals[0][0]['version_family'], originals[1][0]['version_family'])
        doc, raw = originals[0]
        self.assertEqual(self.request('PATCH', '/api/documents/' + doc['id'], {'content': 'Edited extracted text'})[0], 200)
        self.assertEqual(self.request('GET', '/api/documents/' + doc['id'] + '/original'), (200, raw))
        for method, route in [('POST', '/api/documents'), ('PATCH', '/api/documents/' + doc['id'])]:
            status, _ = self.request(method, route, {'title': 'Cannot replace originals', 'attachment_ref': doc['attachment_ref']})
            self.assertEqual(status, 400)
        self.assertEqual(self.request('POST', '/api/projects/' + project['id'] + '/organize', {})[0], 200)
        self.assertEqual(self.request('POST', '/api/projects/' + project['id'] + '/organize', {}, csrf=False)[0], 403)

    def test_original_upload_rejects_absolute_or_traversal_source(self):
        for source in ('../private.txt', '/private.txt', 'C:/private.txt', 'folder/../../private.txt'):
            status, _ = self.request('POST', '/api/upload', {'name': 'file.txt', 'source_ref': source,
                'base64': base64.b64encode(b'Synthetic text').decode()})
            self.assertEqual(status, 400, source)
        self.assertFalse(self.app.stores['personal'].list('documents'))
        self.assertFalse((self.app.data_dir / 'originals').exists())

    def test_missing_original_has_specific_error(self):
        doc = self.create('documents', {'title': 'Synthetic legacy source', 'content': 'Kept text'})
        status, result = self.request('GET', '/api/documents/' + doc['id'] + '/original')
        self.assertEqual(status, 404)
        self.assertEqual(result['code'], 'original_unavailable')
        self.assertIn('原文件', result['error'])

    def test_txt_upload_local_ask_citations_and_selection(self):
        text = '合成星河公司收入为100百万元。收入增长来自合成客户订单。'
        encoded = base64.b64encode(text.encode('utf-8')).decode('ascii')
        status, doc = self.request('POST', '/api/upload', {'name': 'synthetic.txt', 'base64': encoded})
        self.assertEqual(status, 201, doc)
        self.assertEqual(doc['content'], text)
        self.assertTrue(doc['chunks'])
        detail = self.request('GET', '/api/documents/' + doc['id'])[1]
        self.assertEqual(detail['content'], text)
        body = {'question': '星河公司收入', 'document_ids': [doc['id']], 'mode': 'local', 'allow_external': False}
        status, result = self.request('POST', '/api/ask', body)
        self.assertEqual(status, 200, result)
        self.assertEqual(result['mode'], 'local')
        self.assertTrue(result['citations'], result)
        for citation in result['citations']:
            self.assertEqual(citation['document_id'], doc['id'])
            self.assertIn(citation['quote'], text)
            self.assertIn(citation['id'], {doc['id'] + ':' + chunk['id'] for chunk in doc['chunks']})
            matching = [chunk for chunk in doc['chunks'] if chunk['ordinal'] == citation['ordinal']]
            self.assertTrue(matching)
            self.assertIn(citation['quote'], matching[0]['text'])
        with patch.object(self.app, 'local_chat') as provider:
            status, guidance = self.request('POST', '/api/ask', {**body, 'document_ids': []})
            self.assertEqual(status, 200, guidance)
            self.assertEqual(guidance['status'], 'needs_input')
            self.assertTrue(guidance['questions'])
            provider.assert_not_called()
        self.assertEqual(self.app.stores['personal'].list('deliverables'), [])
        self.assertEqual(self.request('POST', '/api/ask', body, workspace='demo')[0], 400)
        self.assertEqual(self.request('POST', '/api/ask', {**body, 'mode': 'model'})[0], 400)
        self.assertEqual(self.request('POST', '/api/upload', {'name': 'bad.txt', 'base64': '!!!'})[0], 400)

    def test_selected_memory_file_upload_is_redacted_local_and_deduplicated(self):
        raw = ('个人记忆内容\nAPI_KEY=sk-' + '0123456789abcdef0123456789abcdef').encode('utf-8')
        encoded = base64.b64encode(raw).decode('ascii')
        body = {'name': 'profile.md', 'base64': encoded, 'kind': 'memory', 'source_ref': 'selected/knowledge/profile.md'}
        status, doc = self.request('POST', '/api/upload', body)
        self.assertEqual(status, 201, doc)
        self.assertEqual(doc['kind'], 'memory')
        self.assertTrue(doc['private'])
        self.assertEqual(doc['source_ref'], 'selected/knowledge/profile.md')
        self.assertNotIn('0123456789abcdef0123456789abcdef', doc['content'])
        self.assertIn('含认证信息的原文行已隐藏', doc['content'])
        status, same = self.request('POST', '/api/upload', body)
        self.assertEqual(status, 201, same)
        self.assertEqual(same['id'], doc['id'])
        self.assertTrue(same['unchanged'])
        revised = {'name': 'profile.md', 'base64': base64.b64encode('偏好已更新'.encode()).decode('ascii'), 'kind': 'memory', 'source_ref': body['source_ref']}
        status, updated = self.request('POST', '/api/upload', revised)
        self.assertEqual(status, 201, updated)
        self.assertEqual(updated['id'], doc['id'])
        self.assertTrue(updated['updated'])
        for path in ('C:/Users/<user>/profile.md', '../credentials/key.md', 'folder/.env'):
            self.assertEqual(self.request('POST', '/api/upload', {**body, 'source_ref': path})[0], 400)
        self.assertEqual(self.request('POST', '/api/upload', body, workspace='demo')[0], 400)
        self.app.dsh_available = True
        with patch.object(self.app, 'dsh_answer') as invoke:
            status, _ = self.request('POST', '/api/ask', {'question': '个人偏好是什么？', 'document_ids': [doc['id']], 'mode': 'dsh', 'allow_external': True})
            self.assertEqual(status, 400)
            invoke.assert_not_called()

    def test_memory_import_only_personal_and_temporary_allowlist(self):
        root = Path(self.tmp.name) / 'synthetic_memory'
        file = root / '知识库/memory/profile.md'
        file.parent.mkdir(parents=True)
        file.write_text('合成偏好：证据先行。', encoding='utf-8')
        self.app.memory_root = root
        rel = '知识库/memory/profile.md'
        scan = self.request('GET', '/api/memory/scan')[1]
        self.assertEqual([f['path'] for f in scan['files']], [rel])
        self.assertNotIn('content', scan['files'][0])
        self.assertEqual(self.request('GET', '/api/memory/scan', workspace='demo')[1]['files'], [])
        self.assertEqual(self.request('POST', '/api/memory/import', {'paths': [rel]}, workspace='demo')[0], 400)
        self.assertEqual(self.request('POST', '/api/memory/import', {'paths': ['../profile.md']})[0], 400)
        status, result = self.request('POST', '/api/memory/import', {'paths': [rel]})
        self.assertEqual(status, 200, result)
        self.assertEqual(result['imported'], 1)

    def test_memory_can_be_retrieved_locally_but_never_sent_to_model(self):
        doc = self.create('documents', {'title': '合成个人记忆', 'kind': 'memory', 'content': '合成记忆偏好：证据先行。'})
        research = self.create('documents', {'title': '合成研究', 'content': '合成研究证据先行。'})
        self.app.ai = {'base_url': 'https://model.invalid/v1', 'model': 'fixture', 'api_key': 'synthetic'}
        body = {'question': '证据', 'document_ids': [doc['id']], 'mode': 'local'}
        self.assertEqual(self.request('POST', '/api/ask', body)[0], 200)
        with patch('workos.server.urllib.request.urlopen', side_effect=AssertionError('Network must not be reached')) as network:
            for ids in ([doc['id']], [research['id'], doc['id']]):
                status, result = self.request('POST', '/api/ask', {**body, 'document_ids': ids, 'mode': 'model', 'allow_external': True})
                self.assertEqual(status, 400, result)
                self.assertIn('记忆', result['error'])
            network.assert_not_called()

    def test_invalid_document_content_is_a_client_error(self):
        before = self.request('GET', '/api/backup')[1]['data']
        for content in (123, None, [], {}):
            with self.subTest(content=content):
                status, result = self.request('POST', '/api/documents', {'title': '合成无效', 'content': content})
                self.assertEqual(status, 400, result)
        self.assertEqual(self.request('GET', '/api/backup')[1]['data'], before)

    def test_dsh_chat_requires_consent_and_uses_only_grounded_citations(self):
        doc = self.create('documents', {'title': '合成 DSH 材料', 'content': '合成星河公司收入为100百万元，收入来自合成订单。'})
        self.app.dsh_available = True
        session_root = Path(self.tmp.name) / 'temporary-sessions'
        overlay = self.app._dsh_overlay('gpt-6-luna', session_root)
        self.assertIn('provider: openai-codex', overlay)
        self.assertIn(json.dumps(str(session_root)), overlay)
        self.assertIn('session-log-deepseek', overlay)
        self.assertIn('model: gpt-6-luna', overlay)
        self.assertIn('- id: tool-bash\n  disabled: true', overlay)
        body = {'question': '星河收入是多少？', 'document_ids': [doc['id']], 'mode': 'dsh', 'model_id': 'gpt-6-luna'}
        with patch.object(self.app, 'dsh_answer', return_value='合成结论：收入100百万元。[S1]') as invoke:
            memory = self.create('documents', {'title': '合成记忆', 'kind': 'memory', 'content': '合成个人偏好'})
            status, denied_memory = self.request('POST', '/api/ask', {**body, 'document_ids': [memory['id']]})
            self.assertEqual(status, 400, denied_memory)
            invoke.assert_not_called()
            status, result = self.request('POST', '/api/ask', body)
        self.assertEqual(status, 200, result)
        self.assertEqual(result['mode'], 'dsh')
        self.assertEqual(result['model'], 'GPT-6 Luna via DSH')
        self.assertTrue(result['citations'])
        self.assertIn('收入为100百万元', invoke.call_args.args[0])
        self.assertIn('星河收入是多少', invoke.call_args.args[0])
        self.assertEqual(invoke.call_args.args[1], 'gpt-6-luna')
        self.assertEqual(self.request('POST', '/api/ask', {**body, 'model_id': 'arbitrary'})[0], 400)

    def test_agent_endpoint_executes_tools_and_blocks_unknown_models(self):
        actions = iter([json.dumps({'action': 'create_note', 'args': {'title': '助手结论', 'body': '合成内容'}, 'say': '保存结论'}),
                        json.dumps({'action': 'final', 'answer': '已保存 1 条结论。'})])
        with patch.object(self.app, 'ai_lock', __import__('threading').RLock()):
            with patch('workos.agent.urllib.request.urlopen') as opened:
                opened.return_value.__enter__.return_value.read.side_effect = lambda *_: json.dumps({'choices': [{'message': {'content': next(actions)}}]}).encode()
                status, result = self.request('POST', '/api/agent', {'message': '存一条结论', 'model_id': 'deepseek-v4.1-flash'})
        self.assertEqual(status, 200, result)
        self.assertEqual(result['answer'], '已保存 1 条结论。')
        self.assertEqual(result['steps'][0]['action'], 'create_note')
        self.assertTrue(any(note['title'] == '助手结论' for note in self.request('GET', '/api/state')[1]['notes']))
    def test_pptx_export_creates_editable_section_slides(self):
        from workos.exports import pptx_report
        from pptx import Presentation
        import io
        payload=pptx_report({'title':'Synthetic Investment Case','body':'# Market\n• Market size is 100.\n## Risks\no Verify adoption.\n➢ If approvals slip, launch moves.'})
        self.assertTrue(payload.startswith(b'PK\x03\x04'))
        deck=Presentation(io.BytesIO(payload))
        self.assertEqual(len(deck.slides),3)
        text='\n'.join(shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)
        for value in ('Synthetic Investment Case','Market','100','Risks','Verify adoption','approvals slip'):
            self.assertIn(value,text)

    def test_minutes_numeric_style_matches_house_rules(self):
        from workos.exports import normalize_minute_numbers, normalize_minute_lines
        cases={'0.13－0.5 μm':'0.13-0.5μm','约 4,000 片/月':'约 4,000 片/月','1000 片':'1,000 片','收入 12500 万元':'收入 12,500 万元','良率 50%－60%':'良率 50%-60%','FY2025A 2026':'FY2025A 2026','2026年 2026-10-03 202607':'2026年 2026-10-03 202607','型号 A1000 001234 12.1000':'型号 A1000 001234 12.1000','1000.5 元 001234 元':'1000.5 元 001234 元','专家—判断':'专家—判断'}
        for raw,expected in cases.items():
            with self.subTest(raw=raw):self.assertEqual(normalize_minute_numbers(raw),expected)
        self.assertEqual(normalize_minute_lines('  • 约 1000 片\n'),'  • 约 1,000 片\n')

    def test_house_minutes_docx_format(self):
        from workos.exports import expert_minutes_docx
        from docx import Document
        docx = expert_minutes_docx('Synthetic Expert Call Notes', '【专家背景】\n王先生，合成职位。\n【专家点评】\n• 合成判断。\n【访谈内容】\n产品定位\n• 技术路线\no 仍需验证\n➢ 合成细节。', 'Synthetic participant', '2026-01-01')
        self.assertTrue(docx.startswith(b'PK\x03\x04'))
        import io
        doc=Document(io.BytesIO(docx));text='\n'.join(p.text for p in doc.paragraphs)
        for required in ('Synthetic Expert Call Notes','专家背景','专家点评','访谈内容','• 合成判断。','o 仍需验证','➢ 合成细节。','Synthetic participant'):
            self.assertIn(required,text)
        self.assertEqual(round(doc.sections[0].page_width.cm,1),21.0)
        self.assertEqual(round(doc.sections[0].page_height.cm,1),29.7)
        self.assertEqual(round(doc.sections[0].top_margin.cm,1),2.2)
        self.assertEqual(doc.paragraphs[1].text,'2026-01-01')
        self.assertEqual(sum(p.text=='2026-01-01' for p in doc.paragraphs),1)

    def test_multi_expert_minutes_has_comparison_matrix_and_contents(self):
        from workos.exports import expert_minutes_docx
        from docx import Document
        import io
        experts=[{'institution':'合成机构'+str(i),'title':'合成职务','date':'2026-01-01','background':'合成背景'+str(i),'comments':['• 合成判断'+str(i)],'content':'采购份额\n• 合成短句'+str(i)} for i in range(5)]
        matrix={'topics':['采购份额','技术路线'],'experts':[0,1,2,3,4],'cells':[[f'{i+1}份额' for i in range(5)],[f'{i+1}路线' for i in range(5)]]}
        docx=expert_minutes_docx('Synthetic Expert Calls','摘要预览','','2026-01-01',experts,matrix,[f'合成机构{i}-合成职务' for i in range(5)])
        doc=Document(io.BytesIO(docx));text='\n'.join(p.text for p in doc.paragraphs)
        self.assertGreaterEqual(len(doc.tables),1)
        table=doc.tables[0];self.assertEqual(len(table.rows),11);self.assertEqual(len(table.columns),2)
        self.assertIn('目录',text)
        self.assertIn('合成机构0-合成职务',text)
        self.assertIn('合成机构0-合成职务\t',text)
        self.assertIn('PAGEREF expert_0',doc.element.xml)
        self.assertNotIn('合成机构0-合成职务\t3',text)
        for value in ('合成机构0','合成机构4','采购份额','合成机构4-合成职务'):
            self.assertIn(value,text)

    def test_meeting_ai_draft_uses_only_selected_transcript(self):
        meeting=self.create('meetings',{'title':'Synthetic Call','transcript':'synthetic transcript only'})
        with patch.object(self.app,'local_chat',return_value=(json.dumps({'title':'Synthetic Call - Expert Call Notes','summary':'【专家背景】\n王先生。\n【专家点评】\n• 合成判断。\n【访谈内容】\n采购情况\n• 原文口径。','participants':'合成专家','date':'2026-01-01','experts':[{'institution':'合成机构','title':'客户总监','background':'合成背景','comments':['• 判断'],'content':'采购份额\n• 口径'}],'matrix':{'topics':['采购份额'],'experts':[0],'cells':[['原文短句'] ]},'contents':[],'warnings':[]},ensure_ascii=False),'deepseek-v4.1-flash')) as chat:
            status,result=self.request('POST','/api/meeting-draft',{'provider':'deepseek','model_id':'deepseek-v4.1-flash','transcript':'ONLY_SELECTED_SYNTHETIC_TRANSCRIPT'})
        self.assertEqual(status,200,result)
        self.assertEqual(result['mode'],'ai');self.assertEqual(result['model'],'deepseek-v4.1-flash')
        self.assertIn('ONLY_SELECTED_SYNTHETIC_TRANSCRIPT',chat.call_args.args[3])
        self.assertNotIn('synthetic transcript only',chat.call_args.args[3])
        self.assertIn('• 原文口径。',result['summary'])
        self.assertEqual(result['matrix']['topics'],['采购份额'])
        self.assertEqual(result['matrix']['cells'][0][0],'原文短句')
        self.assertEqual(result['experts'][0]['institution'],'合成机构')
        payload={key:result[key] for key in ('summary','experts','matrix','contents')}
        status,saved=self.request('PATCH','/api/meetings/'+meeting['id'],payload)
        self.assertEqual(status,200,saved)
        status,reopened=self.request('GET','/api/meetings/'+meeting['id'])
        self.assertEqual(status,200,reopened)
        self.assertEqual(reopened['matrix'],result['matrix'])
        self.assertEqual(reopened['experts'],result['experts'])
        status,word=self.request('GET','/api/meeting-export/'+meeting['id']+'?format=docx')
        self.assertEqual(status,200,word)
        import io
        from docx import Document
        document=Document(io.BytesIO(word))
        self.assertEqual(len(document.tables),1)
        self.assertIn('原文短句',document.tables[0].cell(1,1).text)
        status,edited=self.request('PATCH','/api/meetings/'+meeting['id'],{'summary':'人工修改后的纪要'})
        self.assertEqual(status,200,edited)
        status,word=self.request('GET','/api/meeting-export/'+meeting['id']+'?format=docx')
        document=Document(io.BytesIO(word))
        self.assertEqual(len(document.tables),0)
        self.assertIn('人工修改后的纪要','\n'.join(p.text for p in document.paragraphs))
        with patch.object(self.app,'local_chat',return_value=(json.dumps({'title':'x','summary':'正文','experts':[None,{'institution':123,'comments':'bad','content':None},{'institution':'ok'}],'matrix':{'topics':'bad','experts':None,'cells':[]},'contents':'bad'},ensure_ascii=False),'deepseek-v4.1-flash')):
            status,messy=self.request('POST','/api/meeting-draft',{'provider':'deepseek','transcript':'t'})
        self.assertEqual(status,200,messy)
        self.assertEqual(len(messy['experts']),2)
        self.assertEqual(messy['experts'][0]['comments'],[])
        self.assertEqual(messy['experts'][1]['institution'],'ok')
        self.assertEqual(messy['matrix'],{'topics':[],'experts':[],'cells':[]})
        self.assertEqual(messy['contents'],[])

    def test_sync_status_and_manual_sync_are_safe_when_not_configured(self):
        status, data = self.request('GET', '/api/sync/status')
        self.assertEqual(status, 200, data)
        self.assertFalse(data['enabled'])
        status, result = self.request('POST', '/api/sync', {})
        self.assertEqual(status, 200, result)
        self.assertFalse(result['enabled'])
        bootstrap = self.request('GET', '/api/bootstrap')[1]
        self.assertFalse(bootstrap['sync']['enabled'])

    def test_valuation_xlsx_export_contains_live_model_formulas(self):
        from openpyxl import load_workbook
        import io
        assumptions={'currency':'RMB','unit':'百万元','period':'FY2025A','net_income':100,'pe_multiple':12}
        status,raw=self.request('POST','/api/model/export-xlsx',{'method':'net_income','assumptions':assumptions,'title':'Synthetic Model'})
        self.assertEqual(status,200);self.assertTrue(raw.startswith(b'PK\x03\x04'))
        wb=load_workbook(io.BytesIO(raw),data_only=False)
        self.assertIn('Summary',wb.sheetnames);self.assertIn('Assumptions',wb.sheetnames)
        self.assertEqual(wb['Summary']['B7'].value,'=Assumptions!$B$5*Assumptions!$B$6')
        ps={'currency':'RMB','unit':'百万元','period':'FY2025A','revenue':500,'ps_multiple':2}
        status,psraw=self.request('POST','/api/model/export-xlsx',{'method':'ps','assumptions':ps})
        self.assertEqual(status,200)
        pswb=load_workbook(io.BytesIO(psraw),data_only=False)
        self.assertEqual(pswb['Summary']['B7'].value,'=Assumptions!$B$5*Assumptions!$B$6')
        self.assertEqual(self.request('POST','/api/model/export-xlsx',{'method':'net_income','assumptions':{'net_income':-1,'pe_multiple':12}})[0],400)
        dcf={'currency':'RMB','unit':'百万元','valuation_date':'2025-12-31','wacc':0.10,'discount_timing':'year_end','terminal_method':'perpetuity','terminal_growth':0.02,'net_debt':10,'minority_interest':0,'forecasts':[{'year':'2026E','ebit':100,'da':10,'capex':20,'delta_nwc':5,'tax_rate':0.25}]}
        status,raw=self.request('POST','/api/model/export-xlsx',{'method':'dcf','assumptions':dcf})
        self.assertEqual(status,200)
        model=load_workbook(io.BytesIO(raw),data_only=False)
        self.assertEqual(model['DCF_Forecast']['I2'].value,'=E2+F2-G2-H2')
        self.assertEqual(model['DCF_Forecast']['L2'].value,'=I2*K2')
        self.assertIn('DCF_Forecast',model.sheetnames)
        lbo={'currency':'RMB','unit':'百万元','entry_date':'2026-01-01','exit_date':'2027-01-01','entry_ev':1000,'entry_debt':500,'entry_fees':10,'minimum_cash':10,'initial_cash':10,'seller_rollover':0,'exit_fees':10,'exit_multiple':10,'forecasts':[{'year':'2026','ebitda':100,'da':10,'capex':20,'delta_nwc':5,'tax_rate':0.25,'interest_rate':0.06,'mandatory_amortization':20,'cash_sweep_pct':0.5}]}
        status,raw=self.request('POST','/api/model/export-xlsx',{'method':'lbo','assumptions':lbo})
        self.assertEqual(status,200)
        model=load_workbook(io.BytesIO(raw),data_only=False);ws=model['LBO_Model']
        self.assertEqual(ws['L2'].value,'=J2*G2')
        self.assertEqual(ws['R2'].value,'=MAX(0,J2-P2-Q2)')
        self.assertEqual(ws['B7'].value,'=B5/B6')
        self.assertTrue(str(ws['B10'].value).startswith('=B7^'))
        self.assertIn('Python Ground Truth',str(ws['A12'].value))
        assumptions = {'currency': 'RMB', 'unit': '百万元', 'period': 'FY2025A', 'net_income': 100, 'pe_multiple': 12}
        status, result = self.request('POST', '/api/model/valuation', {'method': 'net_income', 'assumptions': assumptions})
        self.assertEqual(status, 200, result)
        self.assertEqual(result['equity_value'], 1200)
        self.app.dsh_available = True
        proposal = {'assumptions': assumptions, 'clarifications': []}
        with patch.object(self.app, 'dsh_answer', return_value=json.dumps(proposal, ensure_ascii=False)) as invoke:
            body = {'method': 'net_income', 'text': 'FY2025净利润100百万元，P/E 12x', 'model_id': 'gpt-6-luna'}
            status, parsed = self.request('POST', '/api/model/parse-assumptions', body)
        self.assertEqual(status, 200, parsed)
        self.assertEqual(parsed['assumptions'], assumptions)
        self.assertEqual(parsed['missing'], [])
        self.assertIn('确认', parsed['warning'])
        self.assertEqual(invoke.call_args.args[1], 'gpt-6-luna')

    def test_api_key_never_echoed_backed_up_or_persisted(self):
        secret = 'synthetic-secret-never-real-12345'
        status, result = self.request('POST', '/api/ai/settings',
                                      {'base_url': 'http://127.0.0.1:9/v1', 'model': 'fixture', 'api_key': secret})
        self.assertEqual(status, 200, result)
        self.assertNotIn('api_key', result)
        for workspace in ('personal', 'demo'):
            for endpoint in ('/api/bootstrap', '/api/backup', '/api/state'):
                status, data = self.request('GET', endpoint, workspace=workspace)
                self.assertEqual(status, 200)
                self.assertNotIn(secret, json.dumps(data))
                self.assertNotIn('api_key', json.dumps(data))
        for path in self.app.data_dir.iterdir():
            if path.is_file():
                self.assertNotIn(secret.encode(), path.read_bytes(), path.name)

    def test_json_and_ai_settings_validation(self):
        self.assertEqual(self.request('POST', '/api/projects', {'name': 'x'}, headers={'Content-Type': 'text/plain'})[0], 400)
        for base in ('http://external.invalid/v1', 'https://user:pass@example.invalid/v1', 'https://example.invalid/v1?api_key=fixture'):
            self.assertEqual(self.request('POST', '/api/ai/settings', {'base_url': base, 'model': 'fixture'})[0], 400)
        # Wrong field types must be a client error, not a server traceback/500.
        for field in ('base_url', 'model', 'api_key'):
            with self.subTest(field=field):
                self.assertEqual(self.request('POST', '/api/ai/settings', {field: 123})[0], 400)

    def test_public_origin_requires_access_authentication_and_same_origin(self):
        self.assertEqual(self.request('GET','/api/health',headers={'Host':'workos.example.com'})[0],403)
        self.assertEqual(self.request('GET','/api/state',headers={'Cf-Connecting-IP':'198.51.100.1'})[0],403)
        self.app.public_origin='https://workos.example.com'
        try:
            self.assertEqual(self.request('GET','/api/bootstrap',headers={'Host':'workos.example.com'})[0],403)
            self.assertEqual(self.request('GET','/api/state',headers={'Cf-Connecting-IP':'198.51.100.1'})[0],403)
            self.assertEqual(self.request('GET','/api/state',headers={'Host':'workos.example.com','Cf-Access-Authenticated-User-Email':'forged@example.invalid'})[0],403)
            with patch.object(self.app.access_validator,'verify',return_value={'sub':'synthetic'}) as verify:
                self.assertEqual(self.request('GET','/api/health',headers={'Host':'workos.example.com','Cf-Access-Jwt-Assertion':'synthetic-valid'})[0],200)
                verify.assert_called_once_with('synthetic-valid')
                self.assertEqual(self.request('GET','/api/health',headers={'Host':'workos.example.com','Origin':'https://evil.invalid'})[0],403)
                self.assertEqual(self.request('GET','/api/health',headers={'Host':'workos.example.com','Origin':'http://workos.example.com'})[0],403)
                self.assertEqual(self.request('POST','/api/projects',{'name':'Synthetic'},headers={'Host':'workos.example.com','Origin':'https://workos.example.com'},csrf=False)[0],403)
                self.assertEqual(self.request('GET','/api/health',headers={'Host':'workos.example.com','Origin':'https://workos.example.com'})[0],200)
        finally:
            self.app.public_origin=''

    def test_signed_access_token_protects_public_workspace_over_http(self):
        try:
            import jwt
            from cryptography.hazmat.primitives.asymmetric import rsa
        except ImportError:self.skipTest("Optional public-auth dependencies unavailable")
        from types import SimpleNamespace
        from unittest.mock import Mock
        from workos.access_auth import AccessValidator
        key=rsa.generate_private_key(public_exponent=65537,key_size=2048)
        client=Mock();client.get_signing_key_from_jwt.return_value=SimpleNamespace(key=key.public_key())
        old=self.app.access_validator;self.app.access_validator=AccessValidator("synthetic-team","synthetic-aud",client)
        self.app.public_origin="https://workos.example.com"
        try:
            headers={"Host":"workos.example.com","Origin":"https://workos.example.com"}
            self.assertEqual(self.request("GET","/api/state",headers=headers)[0],403)
            now=int(__import__("time").time())
            token=jwt.encode({"sub":"synthetic-user","iss":"https://synthetic-team.cloudflareaccess.com","aud":["synthetic-aud"],"iat":now-1,"exp":now+60},key,algorithm="RS256",headers={"kid":"synthetic-key"})
            headers["Cf-Access-Jwt-Assertion"]=token
            status,data=self.request("GET","/api/state",headers=headers)
            self.assertEqual(status,200,data)
            self.assertIn("projects",data)
            self.assertEqual(self.request("POST","/api/projects",{"name":"Synthetic public write"},headers=headers,csrf=False)[0],403)
            self.assertEqual(self.request("POST","/api/projects",{"name":"Synthetic authenticated write"},headers=headers)[0],201)
        finally:
            self.app.public_origin="";self.app.access_validator=old

    def test_export_pptx_endpoint(self):
        record=self.create('deliverables',{'title':'Synthetic PPT','body':'# Market\n• Market size 100.'})
        status,payload=self.request('GET','/api/export/'+record['id']+'?format=pptx')
        self.assertEqual(status,200)
        self.assertTrue(payload.startswith(b'PK\x03\x04'))

    def test_export_html_escapes_user_markup(self):
        title = '<script>fixture_title_attack()</script>'
        body = '## <img src=x onerror="fixture_body_attack()">\n<script>fixture_script_attack()</script>\n& raw'
        doc = self.create('deliverables', {'title': title, 'body': body})
        status, result = self.request('GET', '/api/export/' + doc['id'] + '?format=html')
        self.assertEqual(status, 200)
        self.assertNotIn(title, result)
        self.assertNotIn('<img src=x', result)
        self.assertNotIn('<script>fixture_script_attack()', result)
        self.assertIn(html.escape(title), result)
        self.assertIn(html.escape('<img src=x onerror="fixture_body_attack()">'), result)
        self.assertIn('&amp; raw', result)
        self.assertIn('id="report-editor-script"', result)
        self.assertIn('id="report-notes-data"', result)
        self.assertNotIn('contenteditable=', result)
        self.assertNotIn('vendor/editable', result)
        self.assertIn('fixture_script_attack()', markdown(doc))
        self.assertEqual(self.request('GET', '/api/export/' + doc['id'] + '?format=unknown')[0], 400)


if __name__ == '__main__':
    unittest.main()
