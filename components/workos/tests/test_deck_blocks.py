import io,json,unittest
from pptx import Presentation
from openpyxl import load_workbook
from workos.exports import pptx_report
from workos.deck_blocks import validate_chart

class DeckBlockTests(unittest.TestCase):
    def chart_data(self):
        return {"type":"column","categories":["2026","2027"],"series":[{"name":"Revenue RMBm","values":[100,120]}],"source":"[S1] Synthetic source\nUnit: RMBm"}
    def test_chart_is_native_editable_and_preserves_embedded_workbook(self):
        data=self.chart_data()
        body="# Revenue\n```chart\n"+json.dumps(data)+"\n```"
        deck=Presentation(io.BytesIO(pptx_report({"title":"QA","body":body})))
        self.assertEqual(len(deck.slides),2)
        chart=next(shape.chart for shape in deck.slides[1].shapes if shape.has_chart)
        self.assertEqual(tuple(chart.series[0].values),(100,120))
        workbook=load_workbook(io.BytesIO(chart.part.chart_workbook.xlsx_part.blob),data_only=True)
        self.assertEqual(workbook.active["B2"].value,100)
        self.assertEqual(workbook.active["B3"].value,120)
        self.assertIn("[S1]",deck.slides[1].notes_slide.notes_text_frame.text)
    def test_long_native_table_preserves_every_row_and_source_text(self):
        rows=["| Item%03d | %d | [S1] |"%(i,i) for i in range(21)]
        body="# Comparison\n| Company | Revenue | Source |\n|---|---:|---|\n"+"\n".join(rows)
        deck=Presentation(io.BytesIO(pptx_report({"title":"QA","body":body})))
        tables=[shape.table for slide in deck.slides for shape in slide.shapes if shape.has_table]
        self.assertEqual(len(tables),4)
        values=[row.cells[0].text for table in tables for row in list(table.rows)[1:]]
        self.assertEqual(values,["Item%03d"%i for i in range(21)])
        self.assertTrue(all(row.cells[2].text=="[S1]" for table in tables for row in list(table.rows)[1:]))
    def test_missing_or_nonfinite_data_is_not_filled_with_zero(self):
        for value in (None,True,"100",float("nan"),float("inf")):
            data=self.chart_data();data["series"][0]["values"][0]=value
            with self.subTest(value=value),self.assertRaises(ValueError):validate_chart(data)
    def test_malformed_chart_and_ragged_table_are_rejected(self):
        for body in ("```chart\n{}", "```chart\nnot json\n```", "|A|B|\n|---|---|\n|one|"):
            with self.subTest(body=body),self.assertRaises(ValueError):pptx_report({"title":"QA","body":body})
    def test_explicit_data_only_no_execution_or_fetch(self):
        data=self.chart_data();data["series"][0]["values"][0]="eval(100)"
        with self.assertRaises(ValueError):validate_chart(data)
