"""Inspect native financial formula dependencies and independent saved snapshots."""
import io
import unittest

from openpyxl import load_workbook
from tests.test_valuation import lbo_fixture
from workos.exports import valuation_xlsx
from workos.valuation import calculate_valuation


class FinancialFormulaExportsTests(unittest.TestCase):
    def test_lbo_ownership_allocations_and_cash_floor_are_formula_linked(self):
        assumptions=lbo_fixture(seller_rollover=100,initial_cash=30)
        result=calculate_valuation('lbo',assumptions)
        workbook=load_workbook(io.BytesIO(valuation_xlsx('lbo',assumptions,result)),data_only=False)
        model=workbook['LBO_Model']
        rows={model.cell(row,1).value:row for row in range(1,model.max_row+1)}
        ass=workbook['Assumptions']
        references={ass.cell(row,1).value:'Assumptions!$B$'+str(row) for row in range(2,ass.max_row+1)}
        equity=rows['Formula Sponsor Equity']
        total_entry=rows['Formula Total Entry Equity']
        ownership=rows['Formula Sponsor Ownership']
        total_exit=rows['Formula Total Exit Proceeds']
        sponsor=rows['Formula Sponsor Proceeds']
        seller=rows['Formula Seller Proceeds']
        self.assertIn(references['initial_cash'],model.cell(equity,2).value)
        self.assertNotIn(references['minimum_cash'],model.cell(equity,2).value)
        self.assertEqual(model.cell(total_entry,2).value,'=B'+str(equity)+'+'+references['seller_rollover'])
        self.assertEqual(model.cell(ownership,2).value,'=B'+str(equity)+'/B'+str(total_entry))
        self.assertEqual(model.cell(sponsor,2).value,'=B'+str(total_exit)+'*B'+str(ownership))
        self.assertEqual(model.cell(seller,2).value,'=B'+str(total_exit)+'-B'+str(sponsor))
        for coordinate in ('P2','Q2','V2'):
            self.assertIn(references['minimum_cash'],model[coordinate].value)
        self.assertEqual(model['T2'].value,'=U2+V2')
        summary={row[0].value:(row[1].value,row[2].value) for row in workbook['Summary'].iter_rows(min_row=7)}
        self.assertEqual(summary['Sponsor Proceeds'],('=LBO_Model!B'+str(sponsor),result['sponsor_proceeds']))
        self.assertEqual(summary['Seller Proceeds'],('=LBO_Model!B'+str(seller),result['seller_proceeds']))
        self.assertAlmostEqual(summary['Sponsor Proceeds'][1]+summary['Seller Proceeds'][1],summary['Total Exit Proceeds'][1])
        self.assertEqual(summary['Sponsor Ownership'][1],result['sponsor_ownership'])
        self.assertEqual(workbook.calculation.calcMode,'auto')
        self.assertTrue(workbook.calculation.fullCalcOnLoad)


if __name__=='__main__':unittest.main()
