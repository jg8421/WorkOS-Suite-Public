import io
import unittest
from pptx import Presentation
from docx import Document
from docx.oxml.ns import qn
from workos.exports import pptx_report, expert_minutes_docx, _slide_pages

class PaginationTests(unittest.TestCase):
    def test_ppt_preserves_more_than_eighteen_items(self):
        items=["Synthetic item %03d"%i for i in range(61)]
        deck=Presentation(io.BytesIO(pptx_report({"title":"QA", "body":"# Topic\n"+"\n".join(items)})))
        text="\n".join(s.text for slide in deck.slides for s in slide.shapes if s.has_text_frame)
        for item in items:self.assertIn(item,text)
        self.assertGreater(len(deck.slides),3)
    def test_ppt_preserves_more_than_forty_sections(self):
        body="\n".join("# Section%03d\nContent%03d"%(i,i) for i in range(45))
        deck=Presentation(io.BytesIO(pptx_report({"title":"QA", "body":body})))
        self.assertEqual(len(deck.slides),46)
        text="\n".join(s.text for slide in deck.slides for s in slide.shapes if s.has_text_frame)
        self.assertIn("Content044",text)
    def test_long_paragraph_is_paginated_without_losing_characters(self):
        paragraph="中文"*500+"END_MARKER"
        pages=_slide_pages([("Topic",[paragraph])])
        rebuilt="".join(line for heading,items in pages for line in items)
        self.assertEqual(rebuilt,paragraph)
        self.assertGreater(len(pages),2)
        self.assertTrue(all(len(items)<=8 for heading,items in pages))
    def test_excessive_deck_is_rejected_not_truncated(self):
        with self.assertRaisesRegex(ValueError,"200"):
            _slide_pages([("Topic",["X"*200000])])
    def test_contents_reference_real_expert_bookmarks(self):
        experts=[{"institution":"Expert%d"%i,"content":"Paragraph\n"*80} for i in range(4)]
        doc=Document(io.BytesIO(expert_minutes_docx("QA","preview",experts=experts)))
        fields=list(doc.element.iter(qn("w:fldSimple")))
        bookmarks=list(doc.element.iter(qn("w:bookmarkStart")))
        names={b.get(qn("w:name")) for b in bookmarks}
        self.assertEqual(len(fields),4)
        for i,field in enumerate(fields):
            self.assertIn("PAGEREF expert_%d"%i,field.get(qn("w:instr")))
            self.assertIn("expert_%d"%i,names)
        self.assertTrue(list(doc.settings.element.iter(qn("w:updateFields"))))
