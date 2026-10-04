"""Synthetic investor economics and real Excel integration, never business data."""
from pathlib import Path
import copy
from io import BytesIO
import math
import os
import tempfile
import unittest
from unittest.mock import patch
from workos.investor_returns import calculate, clarification
from workos.clarifications import resolve_method
from workos.valuation import calculate_valuation
from workos.return_excel import workbook, ExcelUnavailable


def inputs(**changes):
    return {**{'currency':'USD','unit':'millions','entry_date':'2027-12-31','exit_date':'2031-12-31',
               'investment_amount':40,'entry_equity_value':800,'entry_valuation_basis':'post_money',
               'exit_net_income':120,'exit_pe_multiple':15,'ipo_dilution':.2},**changes}


class InvestorEconomicsTests(unittest.TestCase):
    def test_return_request_overrides_company_pe_but_preserves_explicit_lbo(self):
        self.assertEqual(resolve_method('net_income','15x P/E退出，算MOC和IRR'),'investor_return')
        self.assertEqual(resolve_method('lbo','LBO IRR'),'lbo')

    def test_company_value_is_not_investor_proceeds_and_ipo_is_applied_once(self):
        a=inputs();before=copy.deepcopy(a);r=calculate_valuation('investor_return',a)
        self.assertEqual(a,before);self.assertEqual(r['exit_equity_value'],1800)
        self.assertAlmostEqual(r['entry_ownership'],.05);self.assertAlmostEqual(r['exit_ownership'],.04)
        self.assertAlmostEqual(r['total_received'],72);self.assertAlmostEqual(r['moc'],1.8)
        self.assertAlmostEqual(r['irr'],1.8**(365/1461)-1)

    def test_pre_money_and_supplied_final_ownership(self):
        self.assertAlmostEqual(calculate(inputs(entry_valuation_basis='pre_money'))['exit_ownership'],40/840*.8)
        self.assertAlmostEqual(calculate(inputs(exit_ownership=.035,ipo_dilution=None))['moic'],1.575)
        with self.assertRaisesRegex(ValueError,'最终退出持股'):calculate(inputs(exit_ownership=.035))

    def test_multiple_dilution_events_and_no_double_ipo(self):
        self.assertAlmostEqual(calculate(inputs(ipo_dilution=None,dilution_events=[.1,.2]))['exit_ownership'],.036)
        with self.assertRaisesRegex(ValueError,'重复'):calculate(inputs(dilution_events=[.2]))

    def test_explicit_holding_years_and_leap_day_calendar(self):
        self.assertEqual(calculate(inputs(exit_date=None,holding_years=4))['exit_date'],'2031-12-31')
        self.assertEqual(calculate(inputs(entry_date='2028-02-29',exit_date=None,holding_years=1))['exit_date'],'2029-02-28')
        self.assertIsNotNone(clarification(inputs(exit_date=None,holding_years=1.5)))
        self.assertIsNotNone(clarification(inputs(holding_years=5)))

    def test_loose_currency_units_dates_and_percent_aliases_are_normalized_without_relabeling(self):
        from workos.clarifications import normalize_assumptions
        a=normalize_assumptions('investor_return',inputs(currency='美元',unit='百万美金',entry_date='2027年12月31日',ipo_dilution='20%',exit_pe_multiple='15倍'))
        self.assertEqual(a['currency'],'USD');self.assertEqual(a['unit'],'millions');self.assertIsNone(clarification(a))
        self.assertIsNotNone(clarification(normalize_assumptions('investor_return',inputs(currency='CNY',unit='百万美元'))))
        self.assertIsNotNone(clarification(inputs(investment_amount='50%')))

    def test_interim_dividend_changes_actual_date_irr(self):
        r=calculate(inputs(interim_cash_flows=[{'date':'2029-12-31','amount':8,'label':'Dividend'}]))
        self.assertAlmostEqual(r['moic'],2);self.assertGreater(r['irr'],calculate(inputs())['irr'])

    def test_complete_cashflow_schedule_has_no_duplicate_initial_or_terminal(self):
        r=calculate({'currency':'USD','unit':'millions','cash_flows':[
            {'date':'2027-12-31','amount':-40},{'date':'2029-12-31','amount':-10},
            {'date':'2030-12-31','amount':5},{'date':'2031-12-31','amount':95}]})
        self.assertEqual(r['total_invested'],50);self.assertEqual(r['total_received'],100);self.assertEqual(r['moic'],2)
        self.assertEqual(len(r['cash_flows']),4)
        npv=sum(f['amount']/(1+r['irr'])**((__import__('datetime').date.fromisoformat(f['date'])-__import__('datetime').date(2027,12,31)).days/365) for f in r['cash_flows'])
        self.assertAlmostEqual(npv,0,places=8)

    def test_loss_zero_and_invalid_dates_never_fabricate_returns(self):
        self.assertLess(calculate(inputs(exit_proceeds=20))['irr'],0)
        r=calculate(inputs(exit_proceeds=0));self.assertEqual(r['moc'],0);self.assertIsNone(r['irr'])
        for a in (inputs(exit_date='2027-01-01'),inputs(investment_amount=True),inputs(ipo_dilution=25),inputs(currency='unknown')):
            with self.subTest(a=a):self.assertIsNotNone(clarification(a))

    def test_ambiguous_irrs_currencies_and_followon_ownership_prompt_instead_of_guessing(self):
        for a in (inputs(interim_cash_flows=[{'date':'2029-12-31','amount':10,'currency':'CNY'}]),
                  inputs(interim_cash_flows=[{'date':'2029-12-31','amount':-10}]),
                  {'currency':'USD','unit':'millions','cash_flows':[{'date':'2027-12-31','amount':-100},{'date':'2028-12-31','amount':230},{'date':'2029-12-31','amount':-132}]}):
            self.assertEqual(len(clarification(a)['questions']),1)

    def test_missing_critical_inputs_one_group_and_no_optional_terms_interrogation(self):
        report=clarification({'currency':'USD','unit':'millions','investment_amount':40})
        self.assertEqual(len(report['questions']),1);self.assertNotIn('fees',report['missing'])
        self.assertIsNone(clarification(inputs()))
        a=inputs();del a['entry_valuation_basis'];self.assertIn('entry_valuation_basis',clarification(a)['missing'])

    def test_conversation_keeps_selected_drivers_instead_of_stale_computed_fields(self):
        from workos.investor_returns import discard_derived
        a=inputs(holding_years=5,entry_ownership=.05,exit_ownership=.04,exit_equity_value=1800,exit_proceeds=72)
        active=discard_derived(a,'持有5年退出，用利润和P/E')
        self.assertNotIn('exit_date',active);self.assertNotIn('exit_proceeds',active)
        self.assertEqual(calculate(active)['exit_date'],'2032-12-31')
        final=discard_derived(a,'最终持股4%，其余沿用')
        self.assertNotIn('ipo_dilution',final);self.assertAlmostEqual(calculate(final)['exit_ownership'],.04)
        cash=discard_derived(a,'按回收金额72百万美元计算')
        self.assertNotIn('exit_net_income',cash);self.assertAlmostEqual(calculate(cash)['moic'],1.8)


@unittest.skipUnless(os.name=='nt' and (Path.home()/'.cache/codex-runtimes/codex-primary-runtime/dependencies/node/bin/node.exe').is_file(),'Native Excel integration requires the configured Windows host')
class NativeExcelReturnsTests(unittest.TestCase):
    def test_excel_recalculated_formulas_caches_and_changed_drivers(self):
        from openpyxl import load_workbook
        for a in (inputs(),inputs(exit_net_income=240,ipo_dilution=None,dilution_events=[.1,.2]),
                  {'currency':'USD','unit':'millions','cash_flows':[{'date':'2027-12-31','amount':-40},{'date':'2029-12-31','amount':-10},{'date':'2030-12-31','amount':5},{'date':'2031-12-31','amount':95}]}):
            with self.subTest(a=a):
                expected=calculate(a);actual,raw=workbook(a)
                self.assertTrue(actual['excel_verified']);self.assertAlmostEqual(actual['irr'],expected['irr'])
                w=load_workbook(BytesIO(raw),data_only=True,read_only=True);s=w['Returns']
                self.assertAlmostEqual(s['B4'].value,expected['moic']);self.assertAlmostEqual(s['B5'].value,expected['irr']);w.close()
                w=load_workbook(BytesIO(raw),data_only=False,read_only=True)
                self.assertTrue(w['Returns']['B5'].value.startswith('=XIRR('));self.assertEqual(w['Returns']['B4'].value,'=B7/B6');w.close()

    def test_editing_exported_profit_dilution_and_date_cells_recalculates_actual_returns(self):
        import json
        from workos.return_excel import _run
        with tempfile.TemporaryDirectory(prefix='workos-synthetic-return-edit-') as directory:
            folder=Path(directory);path=folder/'synthetic.xlsx';receipt=folder/'result.json'
            path.write_bytes(workbook(inputs())[1])
            _run(['powershell.exe','-NoProfile','-NonInteractive','-ExecutionPolicy','Bypass','-File',
                  str(Path(__file__).with_name('recalculate_return_edit.ps1')),'-WorkbookPath',str(path),'-ResultPath',str(receipt)])
            actual=json.loads(receipt.read_text(encoding='utf-8-sig'))
        expected=calculate(inputs(exit_net_income=240,ipo_dilution=.1,exit_date='2032-12-31'))
        self.assertAlmostEqual(actual['moic'],expected['moic']);self.assertAlmostEqual(actual['irr'],expected['irr']);self.assertAlmostEqual(actual['received'],162)


if __name__=='__main__':unittest.main()
