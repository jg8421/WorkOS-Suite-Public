import io
import tempfile
import unittest
from pathlib import Path
from docx import Document
from workos.store import Store
from workos.exports import expert_minutes_docx

class MeetingPersistenceTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory()
        self.path=Path(self.tmp.name)/"test.sqlite3"
        self.store=Store(self.path)
        self.experts=[{"institution":"Synthetic A","title":"Buyer","background":"Source-backed background","comments":["View"],"content":"Topic\n• 100 units"},{"institution":"Synthetic B","title":"Engineer","content":"Topic\n• 200 units"}]
        self.matrix={"topics":["Topic"],"experts":[0,1],"cells":[["100 units","200 units"]]}
    def tearDown(self):
        self.store.close()
        self.tmp.cleanup()
    def record(self):
        return self.store.create("meetings",{"title":"Synthetic Call","summary":"Source-backed notes","experts":self.experts,"matrix":self.matrix,"contents":["Synthetic A","Synthetic B"]})
    def test_reopen_and_export_preserves_matrix(self):
        record=self.record()
        self.store.close();self.store=Store(self.path)
        saved=self.store.get("meetings",record["id"])
        self.assertEqual(saved["experts"],self.experts)
        self.assertEqual(saved["matrix"],self.matrix)
        raw=expert_minutes_docx(saved["title"],saved["summary"],experts=saved["experts"],matrix=saved["matrix"],contents=saved["contents"])
        doc=Document(io.BytesIO(raw))
        self.assertEqual(len(doc.tables),1)
        self.assertIn("Synthetic A",doc.tables[0].cell(1,1).text)
        self.assertIn("200 units",doc.tables[0].cell(2,1).text)
    def test_backup_restore_preserves_structure(self):
        record=self.record();backup=self.store.backup("personal")
        self.store.restore(backup,"personal")
        self.assertEqual(self.store.get("meetings",record["id"])["matrix"],self.matrix)
    def test_manual_edit_invalidates_stale_export_structure(self):
        record=self.record()
        self.store.update("meetings",record["id"],{"participants":"Someone"})
        self.assertEqual(self.store.get("meetings",record["id"])["matrix"],self.matrix)
        updated=self.store.update("meetings",record["id"],{"summary":"User corrected 100 to 110"})
        self.assertEqual(updated["experts"],[])
        self.assertEqual(updated["matrix"],{})
        self.assertEqual(updated["contents"],[])
        unchanged=self.store.update("meetings",record["id"],{"summary":"Regenerated", "experts":self.experts,"matrix":self.matrix,"contents":[]})
        self.assertEqual(unchanged["matrix"],self.matrix)
    def test_invalid_structure_is_rejected(self):
        for bad in (True,-1,2,"0"):
            with self.subTest(index=bad),self.assertRaises(ValueError):
                self.store.create("meetings",{"title":"Invalid","experts":self.experts,"matrix":{"topics":["Topic"],"experts":[bad],"cells":[["wrong"]]}})
        with self.assertRaises(ValueError):
            self.store.create("meetings",{"title":"Invalid","experts":self.experts,"matrix":{"topics":["Topic"],"experts":[0,1],"cells":[["missing second column"]]}})
        with self.assertRaises(ValueError):
            self.store.create("meetings",{"title":"Invalid","experts":[],"contents":[123]})
