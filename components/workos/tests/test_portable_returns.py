"""New-machine workbook authoring and owned Excel calculation."""
from pathlib import Path
import os
import tempfile
import unittest
from unittest.mock import patch
from openpyxl import load_workbook
from workos.portable_return_workbook import author
from workos.investor_returns import calculate
from tests.test_investor_returns import inputs


def has_excel():
    if os.name != 'nt': return False
    import winreg
    try:
        with winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, 'Excel.Application\\CLSID'): return True
    except OSError: return False


class PortableWorkbookTests(unittest.TestCase):
    def test_zero_ipo_does_not_conflict_with_other_dilution_events(self):
        actual = calculate(inputs(ipo_dilution=0, dilution_events=[.1, .05]))
        self.assertAlmostEqual(actual['exit_ownership'], .05*.9*.95)

    def test_formulas_and_editable_drivers_remain_live_without_node(self):
        a = inputs(ipo_dilution=0,dilution_events=[.1, .05], assumption_sources={'income': '=SYNTHETIC_FORMULA'})
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary)/'synthetic.xlsx'; author(a, calculate(a), path)
            wb = load_workbook(path)
            s = wb['Returns']
            self.assertEqual(s['B4'].value, '=B7/B6')
            self.assertTrue(s['B5'].value.startswith('=XIRR('))
            self.assertEqual(s['B25'].value, '=B11')
            self.assertEqual(s['C26'].value, '=E14')
            self.assertIn('PRODUCT(1-G10,1-G11', s['E12'].value)
            self.assertEqual(s['B15'].font.color.rgb, '000000FF')
            self.assertEqual(wb['Sources']['B2'].data_type, 's')
            self.assertIsNone(wb.vba_archive)
            self.assertFalse(wb._external_links)
            wb.close()

    def test_full_schedule_does_not_add_initial_or_exit_cash_flows(self):
        a = {'currency': 'USD', 'unit': 'millions', 'cash_flows': [
            {'date': '2027-12-31', 'amount': -40}, {'date': '2028-12-31', 'amount': -10}, {'date': '2031-12-31', 'amount': 80}]}
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary)/'synthetic.xlsx'; author(a, calculate(a), path)
            wb = load_workbook(path); s = wb['Returns']
            self.assertEqual(s['B6'].value, '=SUM(B25:B27)')
            self.assertEqual(s['E10'].value, 'Cash-flow schedule below')
            self.assertEqual(s['C27'].value, 80); wb.close()

    @unittest.skipUnless(has_excel(), 'Native Excel is an external Office dependency')
    def test_native_excel_verifies_portable_author_without_codex_runtime(self):
        from workos.return_excel import _workbook, _run
        import json
        with tempfile.TemporaryDirectory(prefix='workos-portable-excel-') as temporary:
            folder = Path(temporary)
            with patch.dict(os.environ, {'WORKOS_ARTIFACT_NODE': str(folder/'missing.exe'), 'WORKOS_ARTIFACT_MODULES': str(folder/'missing')}):
                for index,a in enumerate((inputs(), inputs(ipo_dilution=0,dilution_events=[.1, .05]))):
                    actual, raw = _workbook(a)
                    self.assertTrue(actual['excel_verified']); self.assertEqual(actual['workbook_author'], 'portable-openpyxl')
                    self.assertAlmostEqual(actual['moic'], calculate(a)['moic'])
                    if index: continue
                    path = folder/'edited.xlsx'; receipt = folder/'result.json'; path.write_bytes(raw)
                    _run(['powershell.exe', '-NoProfile', '-NonInteractive', '-ExecutionPolicy', 'Bypass', '-File',
                          str(Path(__file__).with_name('recalculate_return_edit.ps1')), '-WorkbookPath', str(path), '-ResultPath', str(receipt)])
                    edited = json.loads(receipt.read_text(encoding='utf-8-sig'))
                    expected = calculate({**a, 'exit_net_income': 240, 'ipo_dilution': .1, 'exit_date': '2032-12-31'})
                    self.assertAlmostEqual(edited['moic'], expected['moic']); self.assertAlmostEqual(edited['irr'], expected['irr'])


if __name__ == '__main__': unittest.main()
