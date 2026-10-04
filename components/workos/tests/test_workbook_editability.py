import io
import unittest
from openpyxl import load_workbook
from workos.exports import valuation_xlsx
from workos.valuation import calculate_valuation

class WorkbookEditabilityTests(unittest.TestCase):
    def make(self,method,a):
        return load_workbook(io.BytesIO(valuation_xlsx(method,a,calculate_valuation(method,a))),data_only=False)
    def ref(self,wb,key):
        sheet=wb["Assumptions"]
        row=next(r for r in range(2,sheet.max_row+1) if sheet.cell(r,1).value==key)
        return "Assumptions!$B$"+str(row)
    def test_dcf_exit_multiple_tax_and_timing_are_live(self):
        a={"currency":"RMB","unit":"millions","valuation_date":"2025-12-31","wacc":.1,"discount_timing":"year_end","terminal_method":"exit_multiple","terminal_multiple":8,"net_debt":100,"minority_interest":10,"tax_rate":.25,"forecasts":[{"year":"2026E","ebit":100,"da":10,"capex":20,"delta_nwc":5}]}
        wb=self.make("dcf",a);sheet=wb["DCF_Forecast"]
        self.assertIn(self.ref(wb,"tax_rate"),sheet["C2"].value)
        self.assertIn(self.ref(wb,"discount_timing"),sheet["J2"].value)
        self.assertEqual(sheet["B5"].value,"=(B2+F2)*"+self.ref(wb,"terminal_multiple"))
        self.assertEqual(wb["Summary"]["C7"].value,calculate_valuation("dcf",a)["enterprise_value"])
    def test_lbo_exit_terms_inherited_rates_and_dates_are_linked(self):
        a={"currency":"RMB","unit":"millions","entry_date":"2026-01-01","exit_date":"2026-12-31","entry_ev":1000,"entry_debt":500,"entry_fees":10,"minimum_cash":10,"initial_cash":10,"seller_rollover":0,"exit_fees":10,"exit_multiple":10,"tax_rate":.25,"interest_rate":.06,"mandatory_amortization":20,"cash_sweep_pct":.5,"forecasts":[{"year":"2026","ebitda":100,"da":10,"capex":20,"delta_nwc":5}]}
        wb=self.make("lbo",a);sheet=wb["LBO_Model"]
        for col,key in (("F","tax_rate"),("G","interest_rate"),("H","mandatory_amortization"),("I","cash_sweep_pct")):
            self.assertEqual(sheet[col+"2"].value,"="+self.ref(wb,key))
        self.assertIn(self.ref(wb,"exit_multiple"),sheet["B4"].value)
        self.assertIn(self.ref(wb,"exit_fees"),sheet["B14"].value)
        self.assertEqual(sheet["B5"].value,"=B14*B16")
        self.assertEqual(sheet["B8"].value,"="+self.ref(wb,"entry_date"))
        self.assertEqual(sheet["B9"].value,"="+self.ref(wb,"exit_date"))
        self.assertEqual(sheet["B10"].value,"=B7^(365/(B9-B8))-1")
        self.assertTrue(wb["Assumptions"].cell(5,2).is_date)
        summary={row[0].value:row[1].value for row in wb["Summary"].iter_rows(min_row=7)}
        self.assertEqual(summary["IRR"],"=LBO_Model!B10")
        self.assertEqual(summary["MOIC"],"=LBO_Model!B7")
    def test_ps_enterprise_value_bridge_is_formula(self):
        a={"currency":"RMB","unit":"millions","period":"FY25","revenue":500,"ps_multiple":2,"net_debt":75}
        wb=self.make("ps",a)
        self.assertEqual(wb["Summary"]["B8"].value,"=B7+"+self.ref(wb,"net_debt"))
        self.assertEqual(wb["Summary"]["C8"].value,1075)
