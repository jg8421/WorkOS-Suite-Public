"""Exercise the delivered interpreter and isolated app with synthetic data only."""
from __future__ import annotations
import argparse
import hashlib
import importlib.util
import json
import os
import shutil
from pathlib import Path
import socket
import subprocess
import tempfile
import time
import urllib.request
import zipfile
from unittest.mock import patch
from tools.package_windows import extract
from tools.deployment_preflight import clean_environment


LIBRARY_SMOKE = r'''
import io,json,sys,zipfile
from pathlib import Path
from workos.exports import html_report
report={'python':sys.version.split()[0],'isolated':bool(sys.flags.isolated),'source':str(Path(__import__('workos').__file__).resolve())}
assert sys.flags.isolated and not __import__('site').ENABLE_USER_SITE
record={'title':'Synthetic Portable Smoke','body':'## Synthetic section\nPublic synthetic text only.'}
assert 'Synthetic Portable Smoke' in html_report(record)
report['html']=True
if sys.argv[1]=='full':
 from workos.exports import docx_report,pptx_report,valuation_xlsx
 from workos.valuation import calculate_valuation
 from openpyxl import load_workbook
 for kind,raw,entry in [('docx',docx_report(record),'word/document.xml'),('pptx',pptx_report(record),'ppt/slides/slide1.xml')]:
  with zipfile.ZipFile(io.BytesIO(raw)) as zipped: assert b'Synthetic' in zipped.read(entry)
  report[kind]=True
 a={'currency':'USD','unit':'millions','period':'FY2025A','net_income':100,'pe_multiple':12}
 result=calculate_valuation('net_income',a)
 assert result['equity_value']==1200
 wb=load_workbook(io.BytesIO(valuation_xlsx('net_income',a,result)),data_only=False)
 assert wb['Summary']['B7'].value.startswith('=')
 report['xlsx_formula_author']=True
 from workos.investor_returns import calculate
 from workos.portable_return_workbook import author
 a={'currency':'USD','unit':'millions','entry_date':'2027-12-31','exit_date':'2031-12-31','investment_amount':40,'entry_equity_value':800,'entry_valuation_basis':'post_money','exit_net_income':120,'exit_pe_multiple':15,'ipo_dilution':.2}
 path=Path(sys.argv[2])/'synthetic-returns.xlsx'
 author(a,calculate(a),path)
 wb=load_workbook(path,data_only=False)
 assert any('XIRR(' in str(cell.value) for row in wb['Returns'] for cell in row)
 assert not calculate(a).get('excel_verified')
 report['portable_return_author']=True
 report['native_excel_recalculation']='not_run'
report.pop('source')
print(json.dumps(report))
'''


def verify(archive, work, *, diagnose=False):
    archive=Path(archive).resolve();work=Path(work).resolve();work.mkdir(parents=True,exist_ok=True)
    expected=archive.with_suffix('.sha256').read_text(encoding='ascii').split()[0]
    if hashlib.sha256(archive.read_bytes()).hexdigest()!=expected:raise ValueError('Package SHA mismatch')
    with tempfile.TemporaryDirectory(prefix='portable-smoke-',dir=work) as temporary:
        root=Path(temporary);extract(archive,root)
        folders=list(root.iterdir())
        if len(folders)!=1 or not folders[0].is_dir():raise ValueError('Expected one package folder')
        package=folders[0];manifest=json.loads((package/'package-manifest.json').read_text(encoding='utf-8'))
        declared={item['path'] for item in manifest['files']}
        actual={path.relative_to(package).as_posix() for path in package.rglob('*') if path.is_file()}
        if actual!=declared|{'package-manifest.json'}:raise ValueError('Manifest file list mismatch')
        for item in manifest['files']:
            path=package/item['path'];raw=path.read_bytes()
            if len(raw)!=item['bytes'] or hashlib.sha256(raw).hexdigest()!=item['sha256']:raise ValueError('Manifest hash mismatch')
        exe=package/'runtime/python/python.exe';app=package/'app';data=root/'synthetic-data'
        env=clean_environment()
        for key in list(env):
            if key.startswith('WORKOS_'):env.pop(key)
        env.update(WORKOS_PUBLIC_ORIGIN='',WORKOS_SYNC_ROOT='',WORKOS_MEMORY_ROOT='')
        # Mimic a machine with no global Python, Node or DSH discovery.
        env['PATH']=str(Path(os.environ.get('SystemRoot','C:/Windows'))/'System32')
        default_isolation=subprocess.check_output([str(exe),'-c',"import sys,site;assert sys.flags.isolated and not site.ENABLE_USER_SITE;print('isolated-without-user-site')"],cwd=app,env=env,timeout=45)
        assert default_isolation.strip()==b'isolated-without-user-site'
        output=subprocess.check_output([str(exe),'-I','-c',LIBRARY_SMOKE,manifest['profile'],str(root)],cwd=app,env=env,timeout=90)
        libraries=json.loads(output.decode('utf-8'))
        with socket.socket() as sock:sock.bind(('127.0.0.1',0));port=sock.getsockname()[1]
        opener=urllib.request.build_opener(urllib.request.ProxyHandler({}))
        base=f'http://127.0.0.1:{port}'
        process=None
        with (root/'synthetic-server.log').open('wb') as log:
            try:
                # Exercise the delivered launcher helper, retaining only the
                # process handle it creates so cleanup cannot stop other apps.
                spec=importlib.util.spec_from_file_location('delivered_preflight',app/'tools/deployment_preflight.py')
                preflight=importlib.util.module_from_spec(spec);spec.loader.exec_module(preflight)
                original_popen=subprocess.Popen
                def start_worker(command,*args,**kwargs):
                    nonlocal process
                    command=list(command)
                    if diagnose:
                        diagnostic='import faulthandler;faulthandler.dump_traceback_later(6,repeat=True);'
                        if command[1:3]==['-I','-c']:command[3]=diagnostic+command[3]
                        elif command[1:4]==['-I','-m','workos.server']:
                            command=command[:2]+['-c',diagnostic+'import runpy;runpy.run_module("workos.server",run_name="__main__",alter_sys=True)']+command[4:]
                        else:raise ValueError('Unrecognized delivered worker command')
                    process=original_popen(command,*args,**kwargs)
                    return process
                with patch.dict(os.environ,env,clear=True),patch.object(preflight.sys,'executable',str(exe)),patch.object(preflight.subprocess,'Popen',start_worker):
                    health=preflight.start(app,data,port,browser=False)
                assert process is not None
                def get(path):
                    with opener.open(base+path,timeout=2) as response:return json.load(response)
                for _ in range(100):
                    try:health=get('/api/health');break
                    except OSError:
                        if process.poll() is not None:raise ValueError('Portable process exited; see scratch smoke log')
                        time.sleep(.1)
                else:raise ValueError('Portable startup exceeded 10 seconds')
                assert health['version']==manifest['version'] and health['app']=='local-workos'
                boot=get('/api/bootstrap')
                assert boot['workspace']=='personal'
                state=get('/api/state')
                assert not state['projects'] and not state['documents']
                assert get('/api/guidance')['content']
                catalog=get('/api/models')
                assert catalog['groups']
                assert not any(model.get('status')=='verified' for group in catalog['groups'] for model in group['models'])
                readiness=get('/api/system/readiness')
                assert readiness['core_ready']
                harness=get('/api/harness')
                assert harness['engine']=='WorkOS' and harness['tools']['names']
                checks={item['id']:item for item in readiness['checks']}
                assert not checks['dsh']['ready']
                assert not checks['models']['ready']
                if manifest['profile']=='full':assert all(checks[name]['ready'] for name in ('docx','pptx','xlsx'))
                request=urllib.request.Request(base+'/api/shutdown',data=b'{}',headers={'Content-Type':'application/json','X-CSRF-Token':boot['csrf']},method='POST')
                shutdown_started=time.monotonic()
                with opener.open(request,timeout=2) as response:assert json.load(response)['stopping']
                process.wait(timeout=8)
                shutdown_ms=round((time.monotonic()-shutdown_started)*1000)
            except BaseException:
                log.flush()
                failure_log=data/'portable-launcher.log'
                shutil.copyfile(failure_log if failure_log.exists() else root/'synthetic-server.log',work/(manifest['profile']+'-failure.log'))
                raise
            finally:
                if process is not None and process.poll() is None:process.terminate();process.wait(timeout=5)
        return {'version':manifest['version'],'profile':manifest['profile'],'files_verified':len(declared),
                'source_revision':manifest['source_revision'],'build_status':manifest.get('build_status','preview'),
                'package_sha256':expected,'libraries':libraries,'isolated_http':True,
                'default_interpreter_isolated':True,'user_site_not_loaded':True,
                'delivered_launcher_exercised':True,
                'shutdown_ms':shutdown_ms,'diagnostic_trace_enabled':diagnose,
                'personal_workspace_empty':True,'model_calls':0,'global_dependencies_used':False}


if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('archive',type=Path)
    parser.add_argument('--work',type=Path,required=True);parser.add_argument('--report',type=Path)
    parser.add_argument('--diagnose',action='store_true',help='Trace only the synthetic test process if shutdown stalls')
    args=parser.parse_args();result=verify(args.archive,args.work,diagnose=args.diagnose)
    rendered=json.dumps(result,ensure_ascii=False,indent=2)
    if args.report:args.report.parent.mkdir(parents=True,exist_ok=True);args.report.write_text(rendered,encoding='utf-8')
    print(rendered)
